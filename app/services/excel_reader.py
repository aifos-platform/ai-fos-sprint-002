from __future__ import annotations

from typing import Any

from openpyxl import load_workbook

from app.services.budget_mapper import (
    detect_budget_period_columns,
    map_budget_columns,
)
from app.services.budget_normalizer import (
    BudgetNormalizer,
)
from app.services.needed_budget_mapper import (
    detect_needed_budget_columns,
    map_needed_budget_columns,
)
from app.services.needed_budget_normalizer import (
    NeededBudgetNormalizer,
)

from app.services.document_detector import (
    detect_document_type,
)
from app.services.gl_mapper import (
    map_gl_columns,
)
from app.services.expected_funding_mapper import (
    map_expected_funding_columns,
)
from app.services.expected_funding_normalizer import (
    ExpectedFundingNormalizer,
)
from app.services.core_cost_coverage_mapper import (
    map_core_cost_coverage_columns,
)
from app.services.core_cost_coverage_normalizer import (
    CoreCostCoverageNormalizer,
)


def inspect_workbook(
    file_path: str,
) -> dict[str, Any]:
    """
    Inspect an AI-FOS workbook and extract data
    according to the detected document type.

    Budget workbooks may contain:

    - Available Budget
    - Needed Budget

    These datasets are intentionally kept separate.
    """

    workbook = load_workbook(
        file_path,
        data_only=True,
    )

    sheet_names = workbook.sheetnames

    #
    # The first worksheet remains the primary sheet
    # used for normal document-type detection.
    #
    worksheet = workbook[
        sheet_names[0]
    ]

    headers = [
        (
            str(cell.value).strip()
            if cell.value is not None
            else ""
        )
        for cell in worksheet[1]
    ]

    document_type = (
        detect_document_type(
            headers
        )
    )

    gl_mapping: dict[
        str,
        str,
    ] = {}

    transactions: list[
        dict[str, Any]
    ] = []

    budget_mapping: dict[
        str,
        str,
    ] = {}

    budget_lines: list[
        dict[str, Any]
    ] = []

    needed_budget_mapping: dict[
        str,
        str,
    ] = {}

    needed_budget_lines: list[
        dict[str, Any]
    ] = []

    needed_budget_sheet_name: str | None = None

    expected_funding_mapping: dict[
        str,
        str,
    ] = {}

    expected_funding_lines: list[
        dict[str, Any]
    ] = []

    expected_funding_sheet_name: str | None = None    

    core_cost_coverage_mapping: dict[
        str,
        str,
    ] = {}

    core_cost_coverage_lines: list[
        dict[str, Any]
    ] = []

    core_cost_coverage_sheet_name: str | None = None  

    #
    # GENERAL LEDGER
    #

    if document_type == (
        "general_ledger"
    ):

        gl_mapping = (
            map_gl_columns(
                headers
            )
        )

        transactions = (
            _extract_rows(
                worksheet=worksheet,
                headers=headers,
                mapping=gl_mapping,
            )
        )

    #
    # BUDGET
    #

    elif document_type == "budget":

        #
        # 1. AVAILABLE / GRANT BUDGET
        #

        budget_mapping = (
            map_budget_columns(
                headers
            )
        )

        budget_period_columns = (
            detect_budget_period_columns(
                headers
            )
        )

        raw_budget_rows = (
            _extract_budget_rows(
                worksheet=worksheet,
                headers=headers,
                mapping=budget_mapping,
                period_columns=budget_period_columns,
            )
        )

        normalizer = (
            BudgetNormalizer()
        )

        for row in raw_budget_rows:

            normalized_line = (
                normalizer.normalize_line(
                    account_number=(
                        row.get(
                            "account_number"
                        )
                    ),
                    budget_line_code=(
                        row.get(
                            "budget_line_code"
                        )
                    ),
                    budget_line_name=(
                        row.get(
                            "budget_line_name"
                        )
                    ),
                    original_budget=(
                        row.get(
                            "original_budget"
                        )
                    ),
                    remaining_secured_budget=(
                        row.get(
                            "remaining_secured_budget"
                        )
                    ),
                    revised_budget=(
                        row.get(
                            "revised_budget"
                        )
                    ),
                    current_budget_usd=(
                        row.get(
                            "current_budget_usd"
                        )
                    ),
                    fund_code=(
                        row.get(
                            "fund_code"
                        )
                    ),
                    fund_name=(
                        row.get(
                            "fund_name"
                        )
                    ),

                    grant_start_date=(
                        row.get(
                            "grant_start_date"
                        )
                    ),
                    grant_end_date=(
                        row.get(
                            "grant_end_date"
                        )
                    ),

                    donor_code=(
                        row.get(
                            "donor_code"
                        )
                    ),
                    donor_name=(
                        row.get(
                            "donor_name"
                        )
                    ),
                    program_code=(
                        row.get(
                            "program_code"
                        )
                    ),
                    program_name=(
                        row.get(
                            "program_name"
                        )
                    ),
                    category_code=(
                        row.get(
                            "category_code"
                        )
                    ),
                    category_name=(
                        row.get(
                            "category_name"
                        )
                    ),
                    donor_line_code=(
                        row.get(
                            "donor_line_code"
                        )
                    ),
                    project_code=(
                        row.get(
                            "project_code"
                        )
                    ),
                    project_name=(
                        row.get(
                            "project_name"
                        )
                    ),
                    original_currency=(
                        row.get(
                            "original_currency"
                        )
                    ),
                    reporting_currency=(
                        row.get(
                            "reporting_currency"
                        )
                    ),
                    exchange_rate=(
                        row.get(
                            "exchange_rate"
                        )
                    ),
                    fiscal_year=(
                        row.get(
                            "fiscal_year"
                        )
                    ),
                    period_amounts=(
                        row.get(
                            "period_amounts"
                        )
                    ),
                    notes=(
                        row.get(
                            "notes"
                        )
                    ),
                )
            )

            if _is_blank_budget_line(
                normalized_line
            ):
                continue

            budget_lines.append(
                normalized_line
            )

        #
        # 2. NEEDED BUDGET
        #
        # Find the worksheet by name rather than
        # assuming that it is always worksheet #2.
        #

        needed_sheet = None

        for sheet_name in sheet_names:

            if (
                str(sheet_name)
                .strip()
                .lower()
                == "needed budget"
            ):
                needed_sheet = workbook[
                    sheet_name
                ]

                needed_budget_sheet_name = (
                    sheet_name
                )

                break


        if needed_sheet is not None:

            needed_headers = [
                (
                    str(cell.value).strip()
                    if cell.value is not None
                    else ""
                )
                for cell in needed_sheet[1]
            ]

            needed_budget_mapping = (
                map_needed_budget_columns(
                    needed_headers
                )
            )

            needed_amount_columns = (
                detect_needed_budget_columns(
                    needed_headers
                )
            )

            needed_column_indexes = {
                field_name: (
                    needed_headers.index(
                        header_name
                    )
                )
                for (
                    field_name,
                    header_name,
                ) in (
                    needed_budget_mapping.items()
                )
                if header_name
                in needed_headers
            }

            needed_normalizer = (
                NeededBudgetNormalizer()
            )

            for row in needed_sheet.iter_rows(
                min_row=2,
                values_only=True,
            ):

                row_data = {
                    field_name: row[
                        column_index
                    ]
                    for (
                        field_name,
                        column_index,
                    ) in (
                        needed_column_indexes.items()
                    )
                }

                #
                # One worksheet row can create one
                # canonical Needed Budget record per
                # detected fiscal-year amount column.
                #
                for amount_column in (
                    needed_amount_columns
                ):

                    amount_header = (
                        amount_column.get(
                            "header"
                        )
                    )

                    if (
                        amount_header
                        not in needed_headers
                    ):
                        continue

                    amount_index = (
                        needed_headers.index(
                            amount_header
                        )
                    )

                    needed_amount = row[
                        amount_index
                    ]

                    normalized_needed_line = (
                        needed_normalizer.normalize_line(
                            program_code=(
                                row_data.get(
                                    "program_code"
                                )
                            ),
                            program_name=(
                                row_data.get(
                                    "program_name"
                                )
                            ),
                            category_code=(
                                row_data.get(
                                    "category_code"
                                )
                            ),
                            category_name=(
                                row_data.get(
                                    "category_name"
                                )
                            ),
                            budget_line_code=(
                                row_data.get(
                                    "budget_line_code"
                                )
                            ),
                            budget_line_name=(
                                row_data.get(
                                    "budget_line_name"
                                )
                            ),
                            budget_notes=(
                                row_data.get(
                                    "budget_notes"
                                )
                            ),
                            employee_responsible=(
                                row_data.get(
                                    "employee_responsible"
                                )
                            ),
                            fiscal_year=(
                                amount_column.get(
                                    "fiscal_year"
                                )
                            ),
                            needed_budget=(
                                needed_amount
                            ),
                        )
                    )

                    #
                    # Ignore completely empty lines,
                    # but preserve explicit zero budgets.
                    #
                    if (
                        not normalized_needed_line.get(
                            "budget_line_code"
                        )
                        and normalized_needed_line.get(
                            "needed_budget"
                        )
                        is None
                    ):
                        continue

                    needed_budget_lines.append(
                        normalized_needed_line
                    )

        #
        # 3. EXPECTED FUNDING
        #
        # Expected Funding is prospective funding and
        # remains completely separate from secured /
        # available funding.
        #

        expected_funding_sheet = None

        for sheet_name in sheet_names:

            if (
                str(sheet_name)
                .strip()
                .lower()
                == "expected funding"
            ):
                expected_funding_sheet = workbook[
                    sheet_name
                ]

                expected_funding_sheet_name = (
                    sheet_name
                )

                break



        if expected_funding_sheet is not None:

            expected_headers = [
                (
                    str(cell.value).strip()
                    if cell.value is not None
                    else ""
                )
                for cell in expected_funding_sheet[1]
            ]

            expected_funding_mapping = (
                map_expected_funding_columns(
                    expected_headers
                )
            )

            expected_column_indexes = {
                field_name: (
                    expected_headers.index(
                        header_name
                    )
                )
                for (
                    field_name,
                    header_name,
                ) in (
                    expected_funding_mapping.items()
                )
                if header_name
                in expected_headers
            }

            expected_normalizer = (
                ExpectedFundingNormalizer()
            )

            for row in (
                expected_funding_sheet.iter_rows(
                    min_row=2,
                    values_only=True,
                )
            ):

                row_data = {
                    field_name: row[
                        column_index
                    ]
                    for (
                        field_name,
                        column_index,
                    ) in (
                        expected_column_indexes.items()
                    )
                }

                #
                # Ignore completely empty worksheet rows.
                #
                if not any(
                    value not in {
                        None,
                        "",
                    }
                    for value in row_data.values()
                ):
                    continue



                normalized_expected_funding = (
                    expected_normalizer.normalize_line(
                        expected_funding_code=(
                            row_data.get(
                                "expected_funding_code"
                            )
                        ),
                        funding_name=(
                            row_data.get(
                                "funding_name"
                            )
                        ),
                        donor_code=(
                            row_data.get(
                                "donor_code"
                            )
                        ),
                        donor_name=(
                            row_data.get(
                                "donor_name"
                            )
                        ),
                        stage=(
                            row_data.get(
                                "stage"
                            )
                        ),
                        probability_percentage=(
                            row_data.get(
                                "probability_percentage"
                            )
                        ),
                        minimum_amount=(
                            row_data.get(
                                "minimum_amount"
                            )
                        ),
                        most_likely_amount=(
                            row_data.get(
                                "most_likely_amount"
                            )
                        ),
                        maximum_amount=(
                            row_data.get(
                                "maximum_amount"
                            )
                        ),
                        expected_decision_date=(
                            row_data.get(
                                "expected_decision_date"
                            )
                        ),
                        expected_first_payment_date=(
                            row_data.get(
                                "expected_first_payment_date"
                            )
                        ),
                        original_currency=(
                            row_data.get(
                                "original_currency"
                            )
                        ),
                        reporting_currency=(
                            row_data.get(
                                "reporting_currency"
                            )
                        ),
                        program_code=(
                            row_data.get(
                                "program_code"
                            )
                        ),
                        project_code=(
                            row_data.get(
                                "project_code"
                            )
                        ),
                        budget_line_code=(
                            row_data.get(
                                "budget_line_code"
                            )
                        ),
                        notes=(
                            row_data.get(
                                "notes"
                            )
                        ),
                    )
                )

                expected_funding_lines.append(
                    normalized_expected_funding
                )  

        #
        # 4. CORE COST COVERAGE
        #
        # Core Cost Coverage remains separate from
        # Available Budget and Expected Funding.
        #

        core_cost_coverage_sheet = None

        for sheet_name in sheet_names:

            if (
                str(sheet_name)
                .strip()
                .lower()
                == "core cost coverage"
            ):
                core_cost_coverage_sheet = workbook[
                    sheet_name
                ]

                core_cost_coverage_sheet_name = (
                    sheet_name
                )

                break                                  

        if core_cost_coverage_sheet is not None:

            core_cost_coverage_headers = [
                (
                    str(cell.value).strip()
                    if cell.value is not None
                    else ""
                )
                for cell in core_cost_coverage_sheet[1]
            ]

            core_cost_coverage_mapping = (
                map_core_cost_coverage_columns(
                    core_cost_coverage_headers
                )
            )

            core_cost_coverage_column_indexes = {
                field_name: (
                    core_cost_coverage_headers.index(
                        header_name
                    )
                )
                for (
                    field_name,
                    header_name,
                ) in (
                    core_cost_coverage_mapping.items()
                )
                if header_name
                in core_cost_coverage_headers
            }

            core_cost_coverage_normalizer = (
                CoreCostCoverageNormalizer()
            ) 

            for row in (
                core_cost_coverage_sheet.iter_rows(
                    min_row=2,
                    values_only=True,
                )
            ):

                row_data = {
                    field_name: row[
                        column_index
                    ]
                    for (
                        field_name,
                        column_index,
                    ) in (
                        core_cost_coverage_column_indexes.items()
                    )
                }

                #
                # Ignore completely empty worksheet rows.
                #
                if not any(
                    value not in {
                        None,
                        "",
                    }
                    for value in row_data.values()
                ):
                    continue 

                normalized_core_cost_coverage = (
                    core_cost_coverage_normalizer.normalize_line(
                        coverage_type=(
                            row_data.get(
                                "coverage_type"
                            )
                        ),
                        fund_code=(
                            row_data.get(
                                "fund_code"
                            )
                        ),
                        budget_line_code=(
                            row_data.get(
                                "budget_line_code"
                            )
                        ),
                        amount=(
                            row_data.get(
                                "amount"
                            )
                        ),
                    )
                )

                core_cost_coverage_lines.append(
                    normalized_core_cost_coverage
                )                                                                  

    return {
        "sheet_count": len(
            sheet_names
        ),
        "sheet_names": sheet_names,
        "sheet_name": (
            worksheet.title
        ),
        "row_count": (
            worksheet.max_row
        ),
        "column_count": (
            worksheet.max_column
        ),
        "headers": headers,
        "document_type": (
            document_type
        ),
        "gl_mapping": (
            gl_mapping
        ),
        "transactions": (
            transactions
        ),
        "budget_mapping": (
            budget_mapping
        ),
        "budget_lines": (
            budget_lines
        ),
        "needed_budget_sheet_name": (
            needed_budget_sheet_name
        ),
        "needed_budget_mapping": (
            needed_budget_mapping
        ),
        "needed_budget_lines": (
            needed_budget_lines
        ),

        "expected_funding_sheet_name": (
            expected_funding_sheet_name
        ),
        "expected_funding_mapping": (
            expected_funding_mapping
        ),
        "expected_funding_lines": (
            expected_funding_lines
        ),

        "core_cost_coverage_sheet_name": (
            core_cost_coverage_sheet_name
        ),
        "core_cost_coverage_mapping": (
            core_cost_coverage_mapping
        ),
        "core_cost_coverage_lines": (
            core_cost_coverage_lines
        ),

    }


def _extract_rows(
    worksheet,
    headers: list[str],
    mapping: dict[str, str],
) -> list[dict[str, Any]]:
    """
    Extract worksheet rows using a detected
    AI-FOS field mapping.
    """

    extracted_rows: list[
        dict[str, Any]
    ] = []

    column_indexes = {
        field_name: headers.index(
            header_name
        )
        for (
            field_name,
            header_name,
        ) in mapping.items()
        if header_name in headers
    }

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        extracted_row = {
            field_name: (
                row[column_index]
            )
            for (
                field_name,
                column_index,
            ) in column_indexes.items()
        }

        if all(
            value in {
                None,
                "",
            }
            for value
            in extracted_row.values()
        ):
            continue

        extracted_rows.append(
            extracted_row
        )

    return extracted_rows

def _extract_budget_rows(
    worksheet,
    headers: list[str],
    mapping: dict[str, str],
    period_columns: list[dict[str, str | int]],
) -> list[dict[str, Any]]:
    """
    Extract Budget rows while preserving canonical
    fields and dynamic year-specific Budget amounts.

    Dynamic period fields are preserved even when
    their cells are blank.

    This allows AI-FOS to distinguish between:

    - null: no amount was provided
    - 0: an explicit zero amount was provided
    - non-zero: an entered financial amount
    """

    extracted_rows: list[
        dict[str, Any]
    ] = []

    column_indexes = {
        field_name: headers.index(
            header_name
        )
        for (
            field_name,
            header_name,
        ) in mapping.items()
        if header_name in headers
    }

    period_column_indexes: list[
        dict[str, Any]
    ] = []

    for period_column in period_columns:

        header_name = period_column.get(
            "header"
        )

        if header_name not in headers:
            continue

        period_column_indexes.append(
            {
                "column_index": headers.index(
                    header_name
                ),
                "header": header_name,
                "type": period_column.get(
                    "type"
                ),
                "fiscal_year": period_column.get(
                    "fiscal_year"
                ),
            }
        )

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        extracted_row = {
            field_name: row[
                column_index
            ]
            for (
                field_name,
                column_index,
            ) in column_indexes.items()
        }

        period_amounts: list[
            dict[str, Any]
        ] = []

        for period_column in period_column_indexes:

            value = row[
                period_column[
                    "column_index"
                ]
            ]

            if isinstance(
                value,
                str,
            ) and not value.strip():
                value = None

            period_amounts.append(
                {
                    "header": period_column[
                        "header"
                    ],
                    "type": period_column[
                        "type"
                    ],
                    "fiscal_year": period_column[
                        "fiscal_year"
                    ],
                    "amount": value,
                }
            )

        extracted_row[
            "period_amounts"
        ] = period_amounts

        canonical_values = [
            value
            for (
                key,
                value,
            ) in extracted_row.items()
            if key != "period_amounts"
        ]

        has_canonical_data = any(
            value not in {
                None,
                "",
            }
            for value in canonical_values
        )

        if (
            not has_canonical_data
            and not period_column_indexes
        ):
            continue

        extracted_rows.append(
            extracted_row
        )

    return extracted_rows

def _is_blank_budget_line(
    budget_line: dict[str, Any],
) -> bool:
    """
    Determine whether a normalized Budget row
    contains useful financial or dimensional data.
    """

    return not any(
        [
            budget_line.get(
                "account_number"
            ),
            budget_line.get(
                "budget_line_code"
            ),
            budget_line.get(
                "budget_line_name"
            ),
            budget_line.get(
                "fund_code"
            ),
            budget_line.get(
                "donor_code"
            ),
            budget_line.get(
                "program_code"
            ),
            budget_line.get(
                "category_code"
            ),
            budget_line.get(
                "donor_line_code"
            ),
            budget_line.get(
                "project_code"
            ),
            budget_line.get(
                "original_budget"
            ),
            budget_line.get(
                "revised_budget"
            ),
            budget_line.get(
                "current_budget_usd"
            ),
        ]
    )