from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# ============================================================
# DESIGN CONSTANTS
# ============================================================

NAVY = "183153"
DARK_TEXT = "1F2937"
MUTED_TEXT = "667085"
LIGHT_BG = "F5F7FA"
LIGHT_BLUE = "EAF2F8"
BORDER_COLOR = "D9E0E8"
WHITE = "FFFFFF"

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor=NAVY,
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BLUE,
)

LIGHT_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BG,
)

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)

CURRENCY_FORMAT = '#,##0.00;[Red](#,##0.00);-'


# ============================================================
# PUBLIC SHEET BUILDER
# ============================================================


def add_standard_cash_flow_statement_sheet(
    workbook: Workbook,
    report: dict[str, Any] | None,
    *,
    organisation_name: str | None = None,
    base_currency: str | None = None,
) -> None:
    """
    Add the validated Standard Cash Flow Statement as a
    worksheet to an existing Excel workbook.

    This function is presentation-only.

    It does not:
    - read the General Ledger;
    - classify cash-flow transactions;
    - calculate cash movements;
    - modify the validated Cash Flow Statement;
    - recalculate financial results.

    The supplied Standard Cash Flow Statement is the
    authoritative source for all displayed report values.
    """

    report = report or {}

    worksheet = workbook.create_sheet(
        title="Cash Flow Statement"
    )

    _configure_sheet(worksheet)

    current_row = 1

    current_row = _write_report_header(
        worksheet=worksheet,
        row=current_row,
        report=report,
        organisation_name=organisation_name,
        base_currency=base_currency,
    )

    current_row = _write_column_headers(
        worksheet=worksheet,
        row=current_row,
    )

    current_row = _write_activity_sections(
        worksheet=worksheet,
        row=current_row,
        report=report,
    )

    current_row = _write_unclassified_movement(
        worksheet=worksheet,
        row=current_row,
        report=report,
    )

    current_row = _write_net_change(
        worksheet=worksheet,
        row=current_row,
        report=report,
    )

    current_row = _write_reconciliation(
        worksheet=worksheet,
        row=current_row + 1,
        report=report,
    )

    current_row = _write_diagnostics(
        worksheet=worksheet,
        row=current_row + 1,
        report=report,
    )

    _write_report_controls(
        worksheet=worksheet,
        row=current_row + 2,
        report=report,
    )

    _apply_print_settings(worksheet)


# ============================================================
# PUBLIC EXCEL GENERATOR
# ============================================================


def generate_standard_cash_flow_statement_excel(
    report: dict[str, Any] | None,
    *,
    organisation_name: str | None = None,
    base_currency: str | None = None,
) -> bytes:
    """
    Generate a professional standalone Cash Flow Statement
    Excel workbook from the persisted Standard Cash Flow
    Statement reporting output.

    This renderer is presentation-only and does not
    recalculate financial results.
    """

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    add_standard_cash_flow_statement_sheet(
        workbook=workbook,
        report=report,
        organisation_name=organisation_name,
        base_currency=base_currency,
    )

    buffer = BytesIO()
    workbook.save(buffer)

    excel_bytes = buffer.getvalue()
    buffer.close()

    return excel_bytes


# ============================================================
# REPORT HEADER
# ============================================================


def _write_report_header(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
    organisation_name: str | None,
    base_currency: str | None,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    title_cell = worksheet.cell(
        row=row,
        column=1,
        value=organisation_name or "AI-FOS",
    )

    title_cell.font = Font(
        name="Aptos Display",
        size=16,
        bold=True,
        color=WHITE,
    )

    title_cell.fill = HEADER_FILL

    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[row].height = 28

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    report_title_cell = worksheet.cell(
        row=row,
        column=1,
        value="Cash Flow Statement",
    )

    report_title_cell.font = Font(
        name="Aptos Display",
        size=18,
        bold=True,
        color=DARK_TEXT,
    )

    report_title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[row].height = 30

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    method = report.get(
        "method",
        "transaction_counter_account_analysis",
    )

    method_cell = worksheet.cell(
        row=row,
        column=1,
        value=f"Method: {method}",
    )

    method_cell.font = Font(
        name="Aptos",
        size=10,
        color=MUTED_TEXT,
        italic=True,
    )

    row += 1

    currency_text = (
        f"Currency: {base_currency}"
        if base_currency
        else "Currency: Reporting Currency"
    )

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    currency_cell = worksheet.cell(
        row=row,
        column=1,
        value=currency_text,
    )

    currency_cell.font = Font(
        name="Aptos",
        size=9,
        color=MUTED_TEXT,
    )

    return row + 2


# ============================================================
# COLUMN HEADERS
# ============================================================


def _write_column_headers(
    *,
    worksheet,
    row: int,
) -> int:
    headers = (
        "Cash Flow Activity",
        "Amount",
    )

    for column, value in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=row,
            column=column,
            value=value,
        )

        cell.fill = HEADER_FILL

        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=WHITE,
        )

        cell.border = THIN_BORDER

        cell.alignment = Alignment(
            horizontal=(
                "right"
                if value == "Amount"
                else "left"
            ),
            vertical="center",
        )

    worksheet.row_dimensions[row].height = 22

    return row + 1


# ============================================================
# CASH FLOW ACTIVITIES
# ============================================================


def _write_activity_sections(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> int:
    sections = _mapping(
        report.get("sections")
    )

    for key in (
        "operating_activities",
        "investing_activities",
        "financing_activities",
    ):
        section = _mapping(
            sections.get(key)
        )

        if not section:
            continue

        label = str(
            section.get("label")
            or key.replace("_", " ").title()
        )

        amount = _safe_float(
            section.get("amount")
        )

        row = _write_amount_row(
            worksheet=worksheet,
            row=row,
            label=label,
            amount=amount,
        )

    return row


def _write_unclassified_movement(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> int:
    amount = _safe_float(
        report.get("unclassified_cash_movement")
    )

    return _write_amount_row(
        worksheet=worksheet,
        row=row,
        label="Unclassified Cash Movement",
        amount=amount,
        muted=abs(amount) < 0.01,
    )


def _write_amount_row(
    *,
    worksheet,
    row: int,
    label: str,
    amount: float,
    muted: bool = False,
) -> int:
    label_cell = worksheet.cell(
        row=row,
        column=1,
        value=label,
    )

    amount_cell = worksheet.cell(
        row=row,
        column=2,
        value=amount,
    )

    amount_cell.number_format = CURRENCY_FORMAT

    amount_cell.alignment = Alignment(
        horizontal="right",
        vertical="center",
    )

    for column in range(1, 3):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.border = THIN_BORDER

        cell.font = Font(
            name="Aptos",
            size=10,
            color=(
                MUTED_TEXT
                if muted
                else DARK_TEXT
            ),
        )

    label_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    return row + 1


# ============================================================
# NET CHANGE IN CASH
# ============================================================


def _write_net_change(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> int:
    net_change = _safe_float(
        report.get("net_change_in_cash")
    )

    label_cell = worksheet.cell(
        row=row,
        column=1,
        value="NET CHANGE IN CASH",
    )

    value_cell = worksheet.cell(
        row=row,
        column=2,
        value=net_change,
    )

    for column in range(1, 3):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER

        cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=WHITE,
        )

    value_cell.number_format = CURRENCY_FORMAT

    value_cell.alignment = Alignment(
        horizontal="right",
    )

    return row + 1


# ============================================================
# RECONCILIATION
# ============================================================


def _write_reconciliation(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> int:
    reconciliation = _mapping(
        report.get("reconciliation")
    )

    if not reconciliation:
        return row

    row = _write_section_header(
        worksheet=worksheet,
        row=row,
        label="Reconciliation",
    )

    rows = (
        (
            "Classified Net Change",
            _safe_float(
                reconciliation.get(
                    "classified_net_change"
                )
            ),
        ),
        (
            "Reported Net Change in Cash",
            _safe_float(
                reconciliation.get(
                    "net_change_in_cash"
                )
            ),
        ),
        (
            "Reconciliation Difference",
            _safe_float(
                reconciliation.get("difference")
            ),
        ),
    )

    for label, amount in rows:
        row = _write_amount_row(
            worksheet=worksheet,
            row=row,
            label=label,
            amount=amount,
        )

    reconciled = reconciliation.get(
        "is_reconciled"
    )

    status = (
        "Reconciled"
        if reconciled is True
        else "Review Required"
    )

    label_cell = worksheet.cell(
        row=row,
        column=1,
        value="Reconciliation Status",
    )

    value_cell = worksheet.cell(
        row=row,
        column=2,
        value=status,
    )

    for column in range(1, 3):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.fill = LIGHT_FILL
        cell.border = THIN_BORDER

        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=DARK_TEXT,
        )

    value_cell.alignment = Alignment(
        horizontal="right",
    )

    return row + 1


# ============================================================
# CLASSIFICATION DIAGNOSTICS
# ============================================================


def _write_diagnostics(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> int:
    diagnostics = _mapping(
        report.get("classification_diagnostics")
    )

    if not diagnostics:
        return row

    row = _write_section_header(
        worksheet=worksheet,
        row=row,
        label="Classification Diagnostics",
    )

    rows = (
        (
            "Document Count",
            _safe_int(
                diagnostics.get("document_count")
            ),
        ),
        (
            "Classified Document Count",
            _safe_int(
                diagnostics.get(
                    "classified_document_count"
                )
            ),
        ),
        (
            "Unclassified Document Count",
            _safe_int(
                diagnostics.get(
                    "unclassified_document_count"
                )
            ),
        ),
        (
            "Activity Detail Count",
            _safe_int(
                diagnostics.get(
                    "activity_detail_count"
                )
            ),
        ),
    )

    for label, value in rows:
        label_cell = worksheet.cell(
            row=row,
            column=1,
            value=label,
        )

        value_cell = worksheet.cell(
            row=row,
            column=2,
            value=value,
        )

        for column in range(1, 3):
            cell = worksheet.cell(
                row=row,
                column=column,
            )

            cell.border = THIN_BORDER

            cell.font = Font(
                name="Aptos",
                size=10,
                color=DARK_TEXT,
            )

        label_cell.alignment = Alignment(
            horizontal="left",
        )

        value_cell.alignment = Alignment(
            horizontal="right",
        )

        row += 1

    return row


def _write_section_header(
    *,
    worksheet,
    row: int,
    label: str,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    cell = worksheet.cell(
        row=row,
        column=1,
        value=label.upper(),
    )

    cell.fill = SECTION_FILL

    cell.font = Font(
        name="Aptos",
        size=11,
        bold=True,
        color=NAVY,
    )

    cell.border = THIN_BORDER

    return row + 1


# ============================================================
# CONTROLS / METHODOLOGY
# ============================================================


def _write_report_controls(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> None:
    validation = _mapping(
        report.get("validation")
    )

    reconciled = validation.get(
        "reconciled"
    )

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=2,
    )

    cell = worksheet.cell(
        row=row,
        column=1,
        value=(
            "AI-FOS Standard Financial Report"
            + (
                " • Reconciled to verified Cash Flow"
                if reconciled is True
                else ""
            )
        ),
    )

    cell.font = Font(
        name="Aptos",
        size=8,
        italic=True,
        color=MUTED_TEXT,
    )

    note = report.get("note")

    if note:
        row += 1

        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=2,
        )

        note_cell = worksheet.cell(
            row=row,
            column=1,
            value=str(note),
        )

        note_cell.font = Font(
            name="Aptos",
            size=8,
            italic=True,
            color=MUTED_TEXT,
        )

        note_cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        worksheet.row_dimensions[row].height = 34


# ============================================================
# WORKSHEET CONFIGURATION
# ============================================================


def _configure_sheet(
    worksheet,
) -> None:
    worksheet.sheet_view.showGridLines = False

    widths = {
        1: 55,
        2: 22,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    worksheet.freeze_panes = "A7"


def _apply_print_settings(
    worksheet,
) -> None:
    worksheet.page_setup.orientation = "portrait"

    worksheet.page_setup.paperSize = (
        worksheet.PAPERSIZE_A4
    )

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.print_options.horizontalCentered = False

    worksheet.page_margins.left = 0.35
    worksheet.page_margins.right = 0.35
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5

    worksheet.oddFooter.center.text = (
        "Generated by AI-FOS"
    )


# ============================================================
# HELPERS
# ============================================================


def _mapping(
    value: Any,
) -> dict[str, Any]:
    if isinstance(value, dict):
        return value

    return {}


def _safe_float(
    value: Any,
    *,
    default: float = 0.0,
) -> float:
    if value in (
        None,
        "",
    ):
        return default

    try:
        return float(value)

    except (
        TypeError,
        ValueError,
    ):
        return default


def _safe_int(
    value: Any,
    *,
    default: int = 0,
) -> int:
    try:
        return int(value)

    except (
        TypeError,
        ValueError,
    ):
        return default