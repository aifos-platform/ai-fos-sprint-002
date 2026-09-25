from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from app.services.standard_income_statement_excel import (
    add_standard_income_statement_sheet,
)
from app.services.standard_balance_sheet_excel import (
    add_standard_balance_sheet_sheet,
)

# ============================================================
# DESIGN CONSTANTS
# ============================================================

NAVY = "183153"
DARK_TEXT = "1F2937"
MUTED_TEXT = "667085"
LIGHT_BG = "F5F7FA"
LIGHT_BLUE = "EAF2F8"
BORDER_COLOR = "D9E0E8"
WHITE = "FFFFFF"

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor=NAVY,
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BLUE,
)

LIGHT_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BG,
)

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)


# ============================================================
# PUBLIC EXCEL GENERATOR
# ============================================================


def generate_cfo_report_excel(
    report: dict[str, Any] | None,
    standard_income_statement: dict[str, Any] | None = None,
    standard_balance_sheet: dict[str, Any] | None = None,
    *,
    organisation_name: str | None = None,
    base_currency: str | None = None,
) -> bytes:
    """
    Generate the AI-FOS CFO Financial Intelligence Excel workbook.

    The renderer is presentation-only.

    It consumes the validated structured CFO report and does not
    recalculate financial results.
    """

    report = report or {}

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    executive_summary = _mapping(
        report.get("executive_summary")
    )

    financial_health = _mapping(
        report.get("financial_health")
    )

    liquidity = _mapping(
        report.get("liquidity")
    )

    budget = _mapping(
        report.get("budget")
    )

    funding = _mapping(
        report.get("funding")
    )

    core_cost_coverage = _mapping(
        report.get("core_cost_coverage")
    )

    risks = _mapping(
        report.get("risks")
    )

    opportunities = _mapping(
        report.get("opportunities")
    )

    recommendations = _mapping(
        report.get("recommendations")
    )

    trends = _mapping(
        report.get("trends")
    )

    forecast = _mapping(
        report.get("forecast")
    )

    methodology = _mapping(
        report.get("methodology")
    )

    _build_executive_summary_sheet(
        workbook=workbook,
        report=report,
        executive_summary=executive_summary,
        financial_health=financial_health,
        liquidity=liquidity,
        funding=funding,
    )

    _build_financial_health_sheet(
        workbook=workbook,
        financial_health=financial_health,
    )

    _build_liquidity_sheet(
        workbook=workbook,
        liquidity=liquidity,
    )

    _build_budget_sheet(
        workbook=workbook,
        budget=budget,
    )

    _build_funding_sheet(
        workbook=workbook,
        funding=funding,
    )

    _build_core_cost_coverage_sheet(
        workbook=workbook,
        core_cost_coverage=core_cost_coverage,
    )    

    _build_risk_sheet(
        workbook=workbook,
        title="Current Risks",
        risks=_list_of_dicts(
            risks.get("current")
        ),
    )

    _build_risk_sheet(
        workbook=workbook,
        title="Forward Risks",
        risks=_list_of_dicts(
            risks.get("forward")
        ),
    )

    _build_opportunities_sheet(
        workbook=workbook,
        opportunities=opportunities,
    )

    _build_recommendations_sheet(
        workbook=workbook,
        recommendations=recommendations,
    )

    _build_trends_sheet(
        workbook=workbook,
        trends=trends,
    )

    _build_forecast_sheet(
        workbook=workbook,
        forecast=forecast,
    )


    if standard_income_statement:
        add_standard_income_statement_sheet(
            workbook=workbook,
            report=standard_income_statement,
            organisation_name=organisation_name,
            base_currency=base_currency,
        )

    if standard_balance_sheet:
        add_standard_balance_sheet_sheet(
            workbook=workbook,
            report=standard_balance_sheet,
            organisation_name=organisation_name,
            base_currency=base_currency,
        )        


    _build_methodology_sheet(
        workbook=workbook,
        methodology=methodology,
        funding=funding,
    )

    buffer = BytesIO()
    workbook.save(buffer)

    excel_bytes = buffer.getvalue()
    buffer.close()

    return excel_bytes


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================


def _build_executive_summary_sheet(
    *,
    workbook: Workbook,
    report: dict[str, Any],
    executive_summary: dict[str, Any],
    financial_health: dict[str, Any],
    liquidity: dict[str, Any],
    funding: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Executive Summary"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 32,
            "B": 22,
            "C": 22,
            "D": 22,
        },
    )

    _sheet_title(
        sheet,
        "CFO Financial Intelligence Report",
        "Management & Board Financial Overview",
        end_column=4,
    )

    row = 5

    version = report.get("version") or "1.0"

    _write_key_value_rows(
        sheet,
        row,
        {
            "Report Type": (
                "CFO Financial Intelligence Report"
            ),
            "Version": version,
            "Evidence Basis": (
                "Validated AI-FOS financial intelligence"
            ),
        },
    )

    row += 5

    _section_title(
        sheet,
        row,
        "Executive Financial Position",
        end_column=4,
    )

    row += 2

    health = _mapping(
        executive_summary.get(
            "financial_health"
        )
    )

    score = (
        health.get("score")
        if health
        else financial_health.get("score")
    )

    rating = (
        health.get("rating")
        or health.get("status")
        or financial_health.get("rating")
        or financial_health.get("status")
    )

    summary_liquidity = _mapping(
        executive_summary.get(
            "liquidity"
        )
    )

    runway = (
        summary_liquidity.get(
            "cash_runway_months"
        )
        or liquidity.get(
            "cash_runway_months"
        )
    )

    funding_summary = _funding_gap_summary(
        funding
    )

    funding_gap = funding_summary.get(
        "funding_gap"
    )

    coverage = funding_summary.get(
        "applied_coverage_percentage"
    )

    metrics = [
        [
            "Financial Health Score",
            score,
            "Rating",
            rating,
        ],
        [
            "Cash Runway Months",
            runway,
            "Funding Gap",
            funding_gap,
        ],
        [
            "Funding Coverage %",
            coverage,
            "Available Cash",
            liquidity.get("available_cash"),
        ],
        [
            "Blocked Cash",
            liquidity.get("blocked_cash"),
            "Total Cash",
            liquidity.get("total_cash"),
        ],
    ]

    for values in metrics:
        for column, value in enumerate(
            values,
            start=1,
        ):
            cell = sheet.cell(
                row=row,
                column=column,
                value=value,
            )

            cell.border = THIN_BORDER

            if column in (1, 3):
                cell.font = Font(
                    bold=True,
                    color=MUTED_TEXT,
                )

                cell.fill = LIGHT_FILL

            else:
                cell.font = Font(
                    bold=True,
                    color=DARK_TEXT,
                )

        row += 1

    row += 2

    _section_title(
        sheet,
        row,
        "Decision Indicators",
        end_column=4,
    )

    row += 2

    indicators = [
        [
            "High Current Risks",
            executive_summary.get(
                "high_current_risk_count"
            ),
        ],
        [
            "High Forward-Looking Risks",
            executive_summary.get(
                "high_forward_risk_count"
            ),
        ],
        [
            "High-Priority Opportunities",
            executive_summary.get(
                "high_opportunity_count"
            ),
        ],
        [
            "Priority CFO Actions",
            executive_summary.get(
                "priority_action_count"
            ),
        ],
    ]

    _write_table(
        sheet,
        start_row=row,
        headers=[
            "Decision Indicator",
            "Count",
        ],
        rows=indicators,
    )

    row += len(indicators) + 4

    _section_title(
        sheet,
        row,
        "Management Focus",
        end_column=4,
    )

    row += 2

    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row + 2,
        end_column=4,
    )

    cell = sheet.cell(
        row=row,
        column=1,
    )

    cell.value = (
        "Management should use this report to prioritize "
        "the highest-risk financial issues, protect "
        "validated funding coverage, maintain liquidity "
        "resilience, and address forward-looking operating "
        "and funding pressures."
    )

    cell.alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    cell.fill = SECTION_FILL
    cell.border = THIN_BORDER
    cell.font = Font(
        color=DARK_TEXT,
    )


# ============================================================
# FINANCIAL HEALTH
# ============================================================


def _build_financial_health_sheet(
    *,
    workbook: Workbook,
    financial_health: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Financial Health"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 34,
            "B": 20,
            "C": 20,
            "D": 70,
        },
    )

    _sheet_title(
        sheet,
        "Financial Health",
        "Validated AI-FOS financial health assessment",
        end_column=4,
    )

    row = 5

    _write_key_value_rows(
        sheet,
        row,
        _select_mapping(
            financial_health,
            [
                "score",
                "maximum",
                "rating",
                "status",
            ],
        ),
    )

    row += 7

    categories = _mapping(
        financial_health.get("categories")
    )

    if categories:
        _section_title(
            sheet,
            row,
            "Financial Health Areas",
            end_column=4,
        )

        row += 2

        rows = []

        for key, item in categories.items():
            item_data = _mapping(item)

            rows.append(
                [
                    _label(key),
                    item_data.get("score"),
                    item_data.get("maximum"),
                    item_data.get("reason"),
                ]
            )

        _write_table(
            sheet,
            start_row=row,
            headers=[
                "Financial Health Area",
                "Score",
                "Maximum",
                "Assessment",
            ],
            rows=rows,
        )

        row += len(rows) + 4

    metrics = _mapping(
        financial_health.get("metrics")
    )

    selected_metrics = _select_mapping(
        metrics,
        [
            "remaining_requirement",
            "applied_secured_funding",
            "funding_gap",
            "applied_coverage_percentage",
            "matched_requirement_count",
            "unmatched_requirement_count",
            "actual_only_grant_count",
            "budget_only_grant_count",
        ],
    )

    if selected_metrics:
        _section_title(
            sheet,
            row,
            "Key Financial Health Metrics",
            end_column=4,
        )

        row += 2

        _write_key_value_rows(
            sheet,
            row,
            selected_metrics,
        )


# ============================================================
# LIQUIDITY
# ============================================================


def _build_liquidity_sheet(
    *,
    workbook: Workbook,
    liquidity: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Liquidity"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 42,
            "B": 34,
        },
    )

    _sheet_title(
        sheet,
        "Liquidity & Cash Runway",
        "Validated liquidity position and runway evidence",
        end_column=2,
    )

    selected = _select_mapping(
        liquidity,
        [
            "total_cash",
            "available_cash",
            "blocked_cash",
            "available_account_count",
            "blocked_account_count",
            "unclassified_cash_account_count",
            "total_expenses",
            "average_monthly_expenses",
            "cash_runway_months",
            "reporting_period_start",
            "reporting_period_end",
            "reporting_month_count",
            "runway_basis",
        ],
    )

    _write_key_value_rows(
        sheet,
        5,
        selected,
    )


# ============================================================
# BUDGET
# ============================================================


def _build_budget_sheet(
    *,
    workbook: Workbook,
    budget: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Budget Performance"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 46,
            "B": 32,
            "C": 30,
            "D": 30,
        },
    )

    _sheet_title(
        sheet,
        "Budget Performance",
        "Validated budget versus actual intelligence",
        end_column=4,
    )

    executive = _mapping(
        budget.get("executive_summary")
    )

    portfolio = _mapping(
        budget.get("portfolio_control")
    )

    source = executive or portfolio or budget

    selected = _select_mapping(
        source,
        [
            "total_budget",
            "total_actual",
            "budgeted_actual",
            "unbudgeted_actual",
            "budget_variance",
            "total_variance",
            "overall_variance_including_unbudgeted",
            "utilization_percentage",
            "budget_line_count",
            "line_count",
            "over_budget_count",
            "within_budget_count",
            "no_budget_count",
        ],
    )

    _write_key_value_rows(
        sheet,
        5,
        selected,
    )

    row = 5 + len(selected) + 3

    organization = _mapping(
        budget.get("organization")
    )

    scope = _select_mapping(
        organization,
        [
            "fund_count",
            "donor_count",
            "program_count",
            "project_count",
        ],
    )

    if scope:
        _section_title(
            sheet,
            row,
            "Portfolio Scope",
            end_column=4,
        )

        row += 2

        _write_key_value_rows(
            sheet,
            row,
            scope,
        )

        row += len(scope) + 3

    alerts = _list_of_dicts(
        budget.get("alerts")
    )

    if alerts:
        _section_title(
            sheet,
            row,
            "Budget Alerts",
            end_column=4,
        )

        row += 2

        alert_rows = []

        for item in alerts:
            alert_rows.append(
                [
                    item.get("title")
                    or item.get("message")
                    or item.get("category")
                    or "Alert",
                    item.get("severity"),
                    item.get("category"),
                    item.get("reason")
                    or item.get("evidence"),
                ]
            )

        _write_table(
            sheet,
            start_row=row,
            headers=[
                "Alert",
                "Severity",
                "Category",
                "Evidence / Reason",
            ],
            rows=alert_rows,
        )


# ============================================================
# FUNDING
# ============================================================


def _build_funding_sheet(
    *,
    workbook: Workbook,
    funding: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Funding & Grants"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 52,
            "B": 34,
        },
    )

    _sheet_title(
        sheet,
        "Funding & Grant Position",
        "Validated funding gap and grant diagnostics",
        end_column=2,
    )

    summary = _funding_gap_summary(
        funding
    )

    selected = _select_mapping(
        summary,
        [
            "total_needed_budget",
            "actual_spending_against_need",
            "remaining_requirement",
            "gross_remaining_secured_funding",
            "explicitly_mapped_secured_funding",
            "applied_secured_funding",
            "excess_eligible_funding",
            "secured_funding_without_budget_line_allocation",
            "funding_gap",
            "applied_coverage_percentage",
            "matched_requirement_count",
            "unmatched_requirement_count",
            "fully_funded_requirement_count",
            "partially_funded_requirement_count",
            "unfunded_requirement_count",
            "no_remaining_requirement_count",
            "requirements_with_dimension_incompatible_funding",
            "requirements_with_period_ineligible_funding",
            "requirements_with_period_unknown_funding",
        ],
    )

    _write_key_value_rows(
        sheet,
        5,
        selected,
    )

    row = 5 + len(selected) + 3

    diagnostics = _mapping(
        funding.get("grant_diagnostics")
    )

    if diagnostics:
        _section_title(
            sheet,
            row,
            "Grant Diagnostics",
            end_column=2,
        )

        row += 2

        diagnostic_rows = []

        for key in (
            "matched_grants",
            "budget_only_grants",
            "actual_only_grants",
        ):
            value = diagnostics.get(key)

            if isinstance(value, list):
                count = len(value)
            elif isinstance(value, int):
                count = value
            else:
                continue

            diagnostic_rows.append(
                [
                    _label(key),
                    count,
                ]
            )

        if diagnostic_rows:
            _write_table(
                sheet,
                start_row=row,
                headers=[
                    "Diagnostic",
                    "Count",
                ],
                rows=diagnostic_rows,
            )


# ============================================================
# RISKS
# ============================================================

# ============================================================
# CORE COST COVERAGE
# ============================================================


def _build_core_cost_coverage_sheet(
    *,
    workbook: Workbook,
    core_cost_coverage: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Core Cost Coverage"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 42,
            "B": 24,
        },
    )

    _sheet_title(
        sheet,
        "Core Cost Coverage",
        (
            "Validated core cost requirement, coverage, "
            "indirect recovery, and remaining gap"
        ),
        end_column=2,
    )

    summary = _mapping(
        core_cost_coverage.get("summary")
    )

    selected = _select_mapping(
        summary,
        [
            "needed_core_cost",
            "direct_grant_coverage",
            "allocated_indirect_recovery",
            "unrestricted_core_funding",
            "remaining_core_cost_gap",
            "core_cost_coverage_percentage",
            "available_indirect_recovery",
            "used_indirect_recovery",
        ],
    )

    _write_key_value_rows(
        sheet,
        5,
        selected,
    )

    lines = _list_of_dicts(
        core_cost_coverage.get("lines")
    )

    if lines:
        rows = []

        for line in lines:
            rows.append(
                [
                    line.get("program_code"),
                    line.get("category_code"),
                    line.get("budget_line_code"),
                    line.get("budget_line_name"),
                    line.get("employee_responsible"),
                    line.get("fiscal_year"),
                    line.get("needed_core_cost"),
                    line.get("direct_grant_coverage"),
                    line.get(
                        "allocated_indirect_recovery"
                    ),
                    line.get(
                        "unrestricted_core_funding"
                    ),
                    line.get(
                        "remaining_core_cost_gap"
                    ),
                ]
            )

        _write_table(
            sheet,
            start_row=16,
            headers=[
                "Program",
                "Category",
                "Budget Line Code",
                "Budget Line Name",
                "Employee Responsible",
                "Fiscal Year",
                "Needed Core Cost",
                "Direct Grant Coverage",
                "Allocated Indirect Recovery",
                "Unrestricted / Core Funding",
                "Remaining Core Cost Gap",
            ],
            rows=rows,
        )   

def _build_risk_sheet(
    *,
    workbook: Workbook,
    title: str,
    risks: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet(
        title
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 34,
            "B": 16,
            "C": 24,
            "D": 70,
            "E": 70,
        },
    )

    _sheet_title(
        sheet,
        title,
        "Validated AI-FOS risk intelligence",
        end_column=5,
    )

    rows = []

    for risk in risks:
        rows.append(
            [
                risk.get("title")
                or "Financial Risk",
                risk.get("severity"),
                risk.get("category"),
                risk.get("evidence"),
                risk.get("recommendation")
                or risk.get("action"),
            ]
        )

    _write_table(
        sheet,
        start_row=5,
        headers=[
            "Risk",
            "Severity",
            "Category",
            "Evidence",
            "Management Response",
        ],
        rows=rows,
    )


# ============================================================
# OPPORTUNITIES
# ============================================================


def _build_opportunities_sheet(
    *,
    workbook: Workbook,
    opportunities: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Opportunities"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 36,
            "B": 16,
            "C": 24,
            "D": 70,
            "E": 70,
        },
    )

    _sheet_title(
        sheet,
        "Financial Opportunities",
        "Validated opportunities kept separate from financial risks",
        end_column=5,
    )

    items = _list_of_dicts(
        opportunities.get("all")
    )

    rows = []

    for item in items:
        rows.append(
            [
                item.get("title")
                or "Financial Opportunity",
                item.get("priority"),
                item.get("category"),
                item.get("evidence"),
                item.get("recommended_action")
                or item.get("action"),
            ]
        )

    _write_table(
        sheet,
        start_row=5,
        headers=[
            "Opportunity",
            "Priority",
            "Category",
            "Evidence",
            "Recommended Action",
        ],
        rows=rows,
    )


# ============================================================
# RECOMMENDATIONS
# ============================================================


def _build_recommendations_sheet(
    *,
    workbook: Workbook,
    recommendations: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "CFO Recommendations"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 8,
            "B": 38,
            "C": 16,
            "D": 24,
            "E": 65,
            "F": 65,
            "G": 55,
        },
    )

    _sheet_title(
        sheet,
        "CFO Recommendations & Priority Actions",
        "Deterministic management recommendations from validated intelligence",
        end_column=7,
    )

    items = _list_of_dicts(
        recommendations.get("all")
    )

    ordered = sorted(
        items,
        key=lambda item: _priority_rank(
            item.get("priority")
        ),
    )

    rows = []

    for number, item in enumerate(
        ordered,
        start=1,
    ):
        rows.append(
            [
                number,
                item.get("title")
                or item.get("action")
                or f"Management Action {number}",
                item.get("priority"),
                item.get("category"),
                item.get("evidence")
                or item.get("reason"),
                item.get("action"),
                item.get("expected_impact"),
            ]
        )

    _write_table(
        sheet,
        start_row=5,
        headers=[
            "#",
            "Recommendation",
            "Priority",
            "Category",
            "Evidence",
            "Action",
            "Expected Impact",
        ],
        rows=rows,
    )


# ============================================================
# TRENDS
# ============================================================


def _build_trends_sheet(
    *,
    workbook: Workbook,
    trends: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Financial Trends"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 34,
            "B": 26,
            "C": 26,
            "D": 24,
            "E": 24,
        },
    )

    _sheet_title(
        sheet,
        "Financial Trends",
        "Validated historical financial trend intelligence",
        end_column=5,
    )

    row = 5

    overview = _select_mapping(
        trends,
        [
            "status",
            "basis",
        ],
    )

    if overview:
        _write_key_value_rows(
            sheet,
            row,
            overview,
        )

        row += len(overview) + 3

    coverage = _mapping(
        trends.get("coverage")
    )

    if coverage:
        _section_title(
            sheet,
            row,
            "Trend Coverage",
            end_column=5,
        )

        row += 2

        _write_key_value_rows(
            sheet,
            row,
            coverage,
        )

        row += len(coverage) + 3

    quality = _mapping(
        trends.get("quality")
    )

    if quality:
        _section_title(
            sheet,
            row,
            "Data Quality",
            end_column=5,
        )

        row += 2

        _write_key_value_rows(
            sheet,
            row,
            quality,
        )

        row += len(quality) + 3

    for heading, key in (
        (
            "Latest Month Comparison",
            "latest_month_comparison",
        ),
        (
            "Latest Year Comparison",
            "latest_year_comparison",
        ),
    ):
        comparison = _mapping(
            trends.get(key)
        )

        if not comparison:
            continue

        _section_title(
            sheet,
            row,
            heading,
            end_column=5,
        )

        row += 2

        current_period = comparison.get(
            "current_period"
        )

        previous_period = comparison.get(
            "previous_period"
        )

        sheet.cell(
            row=row,
            column=1,
            value="Current Period",
        )

        sheet.cell(
            row=row,
            column=2,
            value=current_period,
        )

        sheet.cell(
            row=row + 1,
            column=1,
            value="Previous Period",
        )

        sheet.cell(
            row=row + 1,
            column=2,
            value=previous_period,
        )

        row += 3

        comparison_rows = []

        for measure in (
            "revenue",
            "expenses",
            "net_result",
        ):
            item = _mapping(
                comparison.get(measure)
            )

            if not item:
                continue

            comparison_rows.append(
                [
                    _label(measure),
                    item.get("current_value"),
                    item.get("previous_value"),
                    item.get("change_percentage"),
                    item.get("direction"),
                ]
            )

        _write_table(
            sheet,
            start_row=row,
            headers=[
                "Measure",
                "Current",
                "Previous",
                "Change %",
                "Direction",
            ],
            rows=comparison_rows,
        )

        row += len(comparison_rows) + 4


# ============================================================
# FORECAST
# ============================================================


def _build_forecast_sheet(
    *,
    workbook: Workbook,
    forecast: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Financial Forecast"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 44,
            "B": 42,
        },
    )

    _sheet_title(
        sheet,
        "Financial Forecast",
        "Deterministic AI-FOS baseline forecast intelligence",
        end_column=2,
    )

    row = 5

    selected = _select_mapping(
        forecast,
        [
            "status",
            "forecast_type",
            "methodology",
            "methodology_description",
            "forecast_horizon_months",
        ],
    )

    if selected:
        _write_key_value_rows(
            sheet,
            row,
            selected,
        )

        row += len(selected) + 3

    for heading, key in (
        (
            "Baseline",
            "baseline",
        ),
        (
            "Forecast Totals",
            "forecast_totals",
        ),
        (
            "Forecast Confidence",
            "confidence",
        ),
    ):
        data = _mapping(
            forecast.get(key)
        )

        if not data:
            continue

        _section_title(
            sheet,
            row,
            heading,
            end_column=2,
        )

        row += 2

        _write_key_value_rows(
            sheet,
            row,
            data,
        )

        row += len(data) + 3


# ============================================================
# METHODOLOGY
# ============================================================


def _build_methodology_sheet(
    *,
    workbook: Workbook,
    methodology: dict[str, Any],
    funding: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet(
        "Methodology"
    )

    _prepare_sheet(
        sheet,
        widths={
            "A": 48,
            "B": 80,
        },
    )

    _sheet_title(
        sheet,
        "Methodology & Evidence",
        "Financial calculation governance and traceability",
        end_column=2,
    )

    row = 5

    if methodology:
        _write_key_value_rows(
            sheet,
            row,
            methodology,
        )

        row += len(methodology) + 3

    funding_gap = _mapping(
        funding.get("funding_gap")
    )

    for possible_key in (
        "methodology",
        "metadata",
    ):
        funding_methodology = _mapping(
            funding_gap.get(
                possible_key
            )
        )

        if not funding_methodology:
            continue

        _section_title(
            sheet,
            row,
            "Funding Gap Methodology",
            end_column=2,
        )

        row += 2

        _write_key_value_rows(
            sheet,
            row,
            funding_methodology,
        )

        row += len(funding_methodology) + 3

        break

    _section_title(
        sheet,
        row,
        "Evidence Principle",
        end_column=2,
    )

    row += 2

    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row + 2,
        end_column=2,
    )

    cell = sheet.cell(
        row=row,
        column=1,
    )

    cell.value = (
        "AI-FOS report presentation consumes validated financial "
        "model outputs. The Excel renderer does not independently "
        "recalculate financial results."
    )

    cell.alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    cell.fill = SECTION_FILL
    cell.border = THIN_BORDER


# ============================================================
# SHARED WORKSHEET HELPERS
# ============================================================


def _prepare_sheet(
    sheet: Worksheet,
    widths: dict[str, float],
) -> None:
    sheet.sheet_view.showGridLines = False

    sheet.freeze_panes = "A5"

    for column, width in widths.items():
        sheet.column_dimensions[
            column
        ].width = width

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    sheet.oddFooter.center.text = (
        "AI-FOS | CONFIDENTIAL — MANAGEMENT & BOARD USE"
    )

    sheet.oddFooter.right.text = (
        "Page &P of &N"
    )


def _sheet_title(
    sheet: Worksheet,
    title: str,
    subtitle: str,
    end_column: int,
) -> None:
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    title_cell = sheet.cell(
        row=1,
        column=1,
        value=title,
    )

    title_cell.font = Font(
        size=18,
        bold=True,
        color=NAVY,
    )

    title_cell.alignment = Alignment(
        vertical="center",
    )

    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=end_column,
    )

    subtitle_cell = sheet.cell(
        row=2,
        column=1,
        value=subtitle,
    )

    subtitle_cell.font = Font(
        size=10,
        color=MUTED_TEXT,
    )

    subtitle_cell.alignment = Alignment(
        vertical="center",
    )

    sheet.merge_cells(
        start_row=3,
        start_column=1,
        end_row=3,
        end_column=end_column,
    )

    line_cell = sheet.cell(
        row=3,
        column=1,
    )

    line_cell.fill = HEADER_FILL

    sheet.row_dimensions[3].height = 4


def _section_title(
    sheet: Worksheet,
    row: int,
    title: str,
    end_column: int,
) -> None:
    sheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=end_column,
    )

    cell = sheet.cell(
        row=row,
        column=1,
        value=title,
    )

    cell.font = Font(
        bold=True,
        size=12,
        color=NAVY,
    )

    cell.fill = SECTION_FILL

    cell.alignment = Alignment(
        vertical="center",
    )

    cell.border = THIN_BORDER

    sheet.row_dimensions[row].height = 22


def _write_key_value_rows(
    sheet: Worksheet,
    start_row: int,
    mapping: dict[str, Any],
) -> None:
    row = start_row

    header_metric = sheet.cell(
        row=row,
        column=1,
        value="Metric",
    )

    header_value = sheet.cell(
        row=row,
        column=2,
        value="Value",
    )

    for cell in (
        header_metric,
        header_value,
    ):
        cell.fill = HEADER_FILL

        cell.font = Font(
            bold=True,
            color=WHITE,
        )

        cell.border = THIN_BORDER

        cell.alignment = Alignment(
            vertical="center",
        )

    row += 1

    for key, value in mapping.items():
        if isinstance(
            value,
            (dict, list),
        ):
            continue

        key_cell = sheet.cell(
            row=row,
            column=1,
            value=_label(key),
        )

        value_cell = sheet.cell(
            row=row,
            column=2,
            value=value,
        )

        key_cell.font = Font(
            bold=True,
            color=DARK_TEXT,
        )

        key_cell.fill = LIGHT_FILL

        key_cell.border = THIN_BORDER
        value_cell.border = THIN_BORDER

        key_cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        value_cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        _apply_number_format(
            value_cell,
            key,
            value,
        )

        row += 1


def _write_table(
    sheet: Worksheet,
    *,
    start_row: int,
    headers: list[str],
    rows: list[list[Any]],
) -> None:
    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=start_row,
            column=column_index,
            value=header,
        )

        cell.fill = HEADER_FILL

        cell.font = Font(
            bold=True,
            color=WHITE,
        )

        cell.border = THIN_BORDER

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    for row_index, values in enumerate(
        rows,
        start=start_row + 1,
    ):
        for column_index, value in enumerate(
            values,
            start=1,
        ):
            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=_excel_safe_value(value),
            )

            cell.border = THIN_BORDER

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            if row_index % 2 == 0:
                cell.fill = LIGHT_FILL


def _apply_number_format(
    cell,
    key: str,
    value: Any,
) -> None:
    if not isinstance(
        value,
        (int, float),
    ) or isinstance(
        value,
        bool,
    ):
        return

    key_lower = key.lower()

    if "percentage" in key_lower:
        # AI-FOS percentage values are already expressed
        # as percentage points, e.g. 75.5 means 75.5%.
        cell.number_format = '0.00"%"'
        return

    if (
        "count" in key_lower
        or "month_count" in key_lower
        or "year_count" in key_lower
    ):
        cell.number_format = "#,##0"
        return

    if (
        "cash" in key_lower
        or "budget" in key_lower
        or "funding" in key_lower
        or "requirement" in key_lower
        or "actual" in key_lower
        or "variance" in key_lower
        or "expenses" in key_lower
        or "revenue" in key_lower
        or "result" in key_lower
    ):
        cell.number_format = "#,##0.00"
        return

    if isinstance(
        value,
        float,
    ):
        cell.number_format = "#,##0.00"


# ============================================================
# DATA HELPERS
# ============================================================


def _mapping(
    value: Any,
) -> dict[str, Any]:
    if isinstance(
        value,
        dict,
    ):
        return value

    return {}


def _list_of_dicts(
    value: Any,
) -> list[dict[str, Any]]:
    if not isinstance(
        value,
        list,
    ):
        return []

    return [
        item
        for item in value
        if isinstance(
            item,
            dict,
        )
    ]


def _select_mapping(
    source: dict[str, Any],
    fields: list[str],
) -> dict[str, Any]:
    result: dict[str, Any] = {}

    for field in fields:
        if field not in source:
            continue

        value = source.get(field)

        if isinstance(
            value,
            (dict, list),
        ):
            continue

        result[field] = value

    return result


def _funding_gap_summary(
    funding: dict[str, Any],
) -> dict[str, Any]:
    funding_gap = _mapping(
        funding.get(
            "funding_gap"
        )
    )

    return _mapping(
        funding_gap.get(
            "summary"
        )
    )


def _priority_rank(
    priority: Any,
) -> int:
    mapping = {
        "critical": 0,
        "high": 1,
        "medium": 2,
        "low": 3,
    }

    return mapping.get(
        str(
            priority or ""
        ).lower(),
        4,
    )


def _label(
    value: Any,
) -> str:
    if value is None:
        return "-"

    return (
        str(value)
        .replace(
            "_",
            " ",
        )
        .strip()
        .title()
    )


def _excel_safe_value(
    value: Any,
) -> Any:
    if value is None:
        return None

    if isinstance(
        value,
        (str, int, float, bool),
    ):
        return value

    return str(value)