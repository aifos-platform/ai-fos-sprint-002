from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# ============================================================
# DESIGN CONSTANTS
# ============================================================

NAVY = "183153"
DARK_TEXT = "1F2937"
MUTED_TEXT = "667085"
LIGHT_BG = "F5F7FA"
LIGHT_BLUE = "EAF2F8"
BORDER_COLOR = "D9E0E8"
WHITE = "FFFFFF"

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor=NAVY,
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BLUE,
)

LIGHT_FILL = PatternFill(
    fill_type="solid",
    fgColor=LIGHT_BG,
)

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)

CURRENCY_FORMAT = '#,##0.00;[Red](#,##0.00);-'


# ============================================================
# PUBLIC WORKSHEET BUILDER
# ============================================================


def add_standard_balance_sheet_sheet(
    workbook: Workbook,
    report: dict[str, Any] | None,
    *,
    organisation_name: str | None = None,
    base_currency: str | None = None,
) -> None:
    """
    Add the validated Standard Balance Sheet as a worksheet
    to an existing Excel workbook.

    This function is presentation-only.

    It does not:
    - read the General Ledger;
    - read the Trial Balance;
    - classify accounts;
    - calculate balances;
    - modify the validated Balance Sheet;
    - recalculate financial results.

    The supplied Standard Balance Sheet is the authoritative
    source for all displayed report values.
    """

    report = report or {}

    worksheet = workbook.create_sheet(
        title="Balance Sheet"
    )

    _configure_sheet(worksheet)

    current_row = 1

    current_row = _write_report_header(
        worksheet=worksheet,
        row=current_row,
        report=report,
        organisation_name=organisation_name,
        base_currency=base_currency,
    )

    current_row = _write_column_headers(
        worksheet=worksheet,
        row=current_row,
    )

    sections = _extract_sections(report)

    for section in sections:
        current_row = _write_section(
            worksheet=worksheet,
            row=current_row,
            section=section,
        )

    current_row = _write_statement_totals(
        worksheet=worksheet,
        row=current_row,
        report=report,
    )

    _write_report_controls(
        worksheet=worksheet,
        row=current_row + 2,
        report=report,
    )

    _apply_print_settings(worksheet)


# ============================================================
# PUBLIC EXCEL GENERATOR
# ============================================================


def generate_standard_balance_sheet_excel(
    report: dict[str, Any] | None,
    *,
    organisation_name: str | None = None,
    base_currency: str | None = None,
) -> bytes:
    """
    Generate a professional standalone Balance Sheet Excel
    workbook from the persisted Standard Balance Sheet
    reporting output.

    This renderer is presentation-only.

    The supplied Standard Balance Sheet is the authoritative
    source for all displayed report values.
    """

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    add_standard_balance_sheet_sheet(
        workbook=workbook,
        report=report,
        organisation_name=organisation_name,
        base_currency=base_currency,
    )

    buffer = BytesIO()
    workbook.save(buffer)

    excel_bytes = buffer.getvalue()
    buffer.close()

    return excel_bytes


# ============================================================
# REPORT HEADER
# ============================================================


def _write_report_header(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
    organisation_name: str | None,
    base_currency: str | None,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=4,
    )

    title_cell = worksheet.cell(
        row=row,
        column=1,
        value=organisation_name or "AI-FOS",
    )

    title_cell.font = Font(
        name="Aptos Display",
        size=16,
        bold=True,
        color=WHITE,
    )

    title_cell.fill = HEADER_FILL

    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[row].height = 28

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=4,
    )

    report_title_cell = worksheet.cell(
        row=row,
        column=1,
        value="Balance Sheet",
    )

    report_title_cell.font = Font(
        name="Aptos Display",
        size=18,
        bold=True,
        color=DARK_TEXT,
    )

    report_title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[row].height = 30

    row += 1

    period_text = _period_text(report)

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=4,
    )

    period_cell = worksheet.cell(
        row=row,
        column=1,
        value=period_text,
    )

    period_cell.font = Font(
        name="Aptos",
        size=10,
        color=MUTED_TEXT,
        italic=True,
    )

    row += 1

    currency_text = (
        f"Currency: {base_currency}"
        if base_currency
        else "Currency: Reporting Currency"
    )

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=4,
    )

    currency_cell = worksheet.cell(
        row=row,
        column=1,
        value=currency_text,
    )

    currency_cell.font = Font(
        name="Aptos",
        size=9,
        color=MUTED_TEXT,
    )

    return row + 2


# ============================================================
# COLUMN HEADERS
# ============================================================


def _write_column_headers(
    *,
    worksheet,
    row: int,
) -> int:
    headers = (
        "Account",
        "Description",
        "Level",
        "Amount",
    )

    for column, value in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=row,
            column=column,
            value=value,
        )

        cell.fill = HEADER_FILL

        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=WHITE,
        )

        cell.border = THIN_BORDER

        cell.alignment = Alignment(
            horizontal=(
                "right"
                if value == "Amount"
                else "left"
            ),
            vertical="center",
        )

    worksheet.row_dimensions[row].height = 22

    return row + 1


# ============================================================
# REPORT SECTIONS
# ============================================================


def _write_section(
    *,
    worksheet,
    row: int,
    section: dict[str, Any],
) -> int:
    section_name = str(
        section.get("section")
        or section.get("name")
        or section.get("title")
        or ""
    ).strip()

    if section_name:
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=4,
        )

        section_cell = worksheet.cell(
            row=row,
            column=1,
            value=section_name.upper(),
        )

        section_cell.fill = SECTION_FILL

        section_cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=NAVY,
        )

        section_cell.border = THIN_BORDER

        row += 1

    lines = _extract_lines(section)

    for line in lines:
        row = _write_report_line(
            worksheet=worksheet,
            row=row,
            line=line,
        )

    if section_name == "Equity / Net Assets":
        reported_equity = _first_numeric(
            section,
            "reported_equity",
        )

        current_period_result = _first_numeric(
            section,
            "current_period_result",
        )

        adjusted_equity = _first_numeric(
            section,
            "total",
        )

        if reported_equity is not None:
            row = _write_total_row(
                worksheet=worksheet,
                row=row,
                label="Total Reported Equity / Net Assets",
                amount=reported_equity,
            )

        if current_period_result is not None:
            row = _write_total_row(
                worksheet=worksheet,
                row=row,
                label="Current Period Surplus / (Deficit)",
                amount=current_period_result,
            )

        if adjusted_equity is not None:
            row = _write_primary_total_row(
                worksheet=worksheet,
                row=row,
                label="Adjusted Equity / Net Assets",
                amount=adjusted_equity,
            )

    else:
        section_total = _first_numeric(
            section,
            "total",
        )

        if section_total is not None:
            row = _write_primary_total_row(
                worksheet=worksheet,
                row=row,
                label=f"Total {section_name}",
                amount=section_total,
            )

    row += 1

    return row


def _write_report_line(
    *,
    worksheet,
    row: int,
    line: dict[str, Any],
) -> int:
    account_number = _first_value(
        line,
        "account_number",
        "account",
        "number",
        "code",
    )

    description = _first_value(
        line,
        "account_name",
        "description",
        "name",
        "label",
        "title",
    )

    level = _safe_int(
        line.get("level"),
        default=0,
    )

    amount = _first_numeric(
        line,
        "amount",
        "balance",
        "value",
    )

    worksheet.cell(
        row=row,
        column=1,
        value=account_number,
    )

    description_cell = worksheet.cell(
        row=row,
        column=2,
        value=description,
    )

    description_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
        indent=max(
            0,
            min(level, 8),
        ),
    )

    worksheet.cell(
        row=row,
        column=3,
        value=level,
    )

    amount_cell = worksheet.cell(
        row=row,
        column=4,
        value=amount,
    )

    amount_cell.number_format = CURRENCY_FORMAT

    amount_cell.alignment = Alignment(
        horizontal="right",
        vertical="center",
    )

    for column in range(1, 5):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.border = THIN_BORDER

        cell.font = Font(
            name="Aptos",
            size=10,
            color=DARK_TEXT,
        )

    return row + 1


def _write_total_row(
    *,
    worksheet,
    row: int,
    label: str,
    amount: float,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=3,
    )

    worksheet.cell(
        row=row,
        column=1,
        value=label,
    )

    amount_cell = worksheet.cell(
        row=row,
        column=4,
        value=amount,
    )

    for column in range(1, 5):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.fill = LIGHT_FILL
        cell.border = THIN_BORDER
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=DARK_TEXT,
        )

    amount_cell.number_format = CURRENCY_FORMAT
    amount_cell.alignment = Alignment(
        horizontal="right",
    )

    return row + 1


def _write_primary_total_row(
    *,
    worksheet,
    row: int,
    label: str,
    amount: float,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=3,
    )

    worksheet.cell(
        row=row,
        column=1,
        value=label,
    )

    amount_cell = worksheet.cell(
        row=row,
        column=4,
        value=amount,
    )

    for column in range(1, 5):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.fill = SECTION_FILL
        cell.border = THIN_BORDER
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=NAVY,
        )

    amount_cell.number_format = CURRENCY_FORMAT
    amount_cell.alignment = Alignment(
        horizontal="right",
    )

    return row + 1


# ============================================================
# STATEMENT TOTALS
# ============================================================


def _write_statement_totals(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> int:
    totals = _mapping(
        report.get("totals")
    )

    liabilities_and_equity = _first_numeric(
        totals,
        "liabilities_and_equity",
    )

    difference = _first_numeric(
        totals,
        "difference",
    )

    if liabilities_and_equity is not None:
        row = _write_final_total(
            worksheet=worksheet,
            row=row,
            label="TOTAL LIABILITIES & EQUITY",
            amount=liabilities_and_equity,
        )

    if difference is not None:
        row = _write_difference_row(
            worksheet=worksheet,
            row=row,
            amount=difference,
        )

    return row


def _write_final_total(
    *,
    worksheet,
    row: int,
    label: str,
    amount: float,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=3,
    )

    worksheet.cell(
        row=row,
        column=1,
        value=label,
    )

    amount_cell = worksheet.cell(
        row=row,
        column=4,
        value=amount,
    )

    for column in range(1, 5):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=WHITE,
        )

    amount_cell.number_format = CURRENCY_FORMAT
    amount_cell.alignment = Alignment(
        horizontal="right",
    )

    return row + 1


def _write_difference_row(
    *,
    worksheet,
    row: int,
    amount: float,
) -> int:
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=3,
    )

    label_cell = worksheet.cell(
        row=row,
        column=1,
        value="BALANCE SHEET DIFFERENCE",
    )

    amount_cell = worksheet.cell(
        row=row,
        column=4,
        value=amount,
    )

    for column in range(1, 5):
        cell = worksheet.cell(
            row=row,
            column=column,
        )

        cell.border = THIN_BORDER
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=DARK_TEXT,
        )

        if amount == 0.0:
            cell.fill = LIGHT_FILL

    amount_cell.number_format = CURRENCY_FORMAT
    amount_cell.alignment = Alignment(
        horizontal="right",
    )

    if amount == 0.0:
        label_cell.value = "BALANCE SHEET DIFFERENCE — RECONCILED"

    return row + 1


# ============================================================
# CONTROLS / METHODOLOGY
# ============================================================


def _write_report_controls(
    *,
    worksheet,
    row: int,
    report: dict[str, Any],
) -> None:
    reconciliation = _mapping(
        report.get("reconciliation")
    )

    reconciled = reconciliation.get(
        "reconciled"
    )

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=4,
    )

    cell = worksheet.cell(
        row=row,
        column=1,
        value=(
            "AI-FOS Standard Financial Report"
            + (
                " • Reconciled to validated Balance Sheet"
                if reconciled is True
                else ""
            )
        ),
    )

    cell.font = Font(
        name="Aptos",
        size=8,
        italic=True,
        color=MUTED_TEXT,
    )


# ============================================================
# WORKSHEET CONFIGURATION
# ============================================================


def _configure_sheet(
    worksheet,
) -> None:
    worksheet.sheet_view.showGridLines = False

    widths = {
        1: 18,
        2: 52,
        3: 10,
        4: 20,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    worksheet.freeze_panes = "A7"


def _apply_print_settings(
    worksheet,
) -> None:
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.paperSize = (
        worksheet.PAPERSIZE_A4
    )
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.print_options.horizontalCentered = False

    worksheet.page_margins.left = 0.35
    worksheet.page_margins.right = 0.35
    worksheet.page_margins.top = 0.5
    worksheet.page_margins.bottom = 0.5

    worksheet.oddFooter.center.text = (
        "Generated by AI-FOS"
    )


# ============================================================
# STRUCTURE HELPERS
# ============================================================


def _extract_sections(
    report: dict[str, Any],
) -> list[dict[str, Any]]:
    value = report.get("sections")

    if not isinstance(value, list):
        return []

    return [
        item
        for item in value
        if isinstance(item, dict)
    ]


def _extract_lines(
    section: dict[str, Any],
) -> list[dict[str, Any]]:
    value = section.get("lines")

    if not isinstance(value, list):
        return []

    return [
        item
        for item in value
        if isinstance(item, dict)
    ]


def _period_text(
    report: dict[str, Any],
) -> str:
    as_of_date = (
        report.get("as_of_date")
        or report.get("analysis_end_date")
        or report.get("latest_closing_date")
    )

    if as_of_date:
        return f"As of {as_of_date}"

    return "Balance Sheet"


def _mapping(
    value: Any,
) -> dict[str, Any]:
    if isinstance(value, dict):
        return value

    return {}


def _first_value(
    mapping: dict[str, Any],
    *keys: str,
) -> Any:
    for key in keys:
        value = mapping.get(key)

        if value not in {
            None,
            "",
        }:
            return value

    return None


def _first_numeric(
    mapping: dict[str, Any],
    *keys: str,
) -> float | None:
    for key in keys:
        value = mapping.get(key)

        if value is None or value == "":
            continue

        try:
            return float(value)

        except (
            TypeError,
            ValueError,
        ):
            continue

    return None


def _safe_int(
    value: Any,
    *,
    default: int = 0,
) -> int:
    try:
        return int(value)

    except (
        TypeError,
        ValueError,
    ):
        return default