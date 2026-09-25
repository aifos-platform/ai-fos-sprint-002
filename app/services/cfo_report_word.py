from io import BytesIO
from typing import Any

from docx import Document


FINANCIAL_HEALTH_METRIC_KEYS = [
    "remaining_requirement",
    "applied_secured_funding",
    "funding_gap",
    "applied_coverage_percentage",
    "matched_requirement_count",
    "unmatched_requirement_count",
    "actual_only_grant_count",
    "budget_only_grant_count",
]

BUDGET_SUMMARY_KEYS = [
    "total_budget",
    "total_actual",
    "budgeted_actual",
    "unbudgeted_actual",
    "budget_variance",
    "total_variance",
    "overall_variance_including_unbudgeted",
    "utilization_percentage",
    "budget_line_count",
    "line_count",
    "over_budget_count",
    "within_budget_count",
    "no_budget_count",
]

BUDGET_SCOPE_KEYS = [
    "fund_count",
    "donor_count",
    "program_count",
    "project_count",
]

FUNDING_SUMMARY_KEYS = [
    "remaining_requirement",
    "applied_secured_funding",
    "excess_eligible_funding",
    "secured_funding_without_budget_line_allocation",
    "funding_gap",
    "applied_coverage_percentage",
    "matched_requirement_count",
    "unmatched_requirement_count",
    "fully_funded_requirement_count",
    "partially_funded_requirement_count",
    "unfunded_requirement_count",
    "no_remaining_requirement_count",
    "requirements_with_dimension_incompatible_funding",
    "requirements_with_period_ineligible_funding",
    "requirements_with_period_unknown_funding",
]

GRANT_DIAGNOSTIC_KEYS = [
    "matched_grants",
    "budget_only_grants",
    "actual_only_grants",
]


def _mapping(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    return {}


def _value(value: Any) -> str:
    if value is None:
        return "-"
    return str(value)


def _currency(value: Any) -> str:
    if value is None:
        return "-"

    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return _value(value)


def _percentage(value: Any) -> str:
    if value is None:
        return "-"

    try:
        return f"{float(value):.1f}%"
    except (TypeError, ValueError):
        return _value(value)


def _display_label(value: Any) -> str:
    text = _value(value)

    if text == "-":
        return text

    return text.replace("_", " ").title()


def _list_of_dicts(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []

    return [
        item
        for item in value
        if isinstance(item, dict)
    ]

def _priority_rank(
    priority: Any,
) -> int:
    mapping = {
        "critical": 0,
        "high": 1,
        "medium": 2,
        "low": 3,
    }

    return mapping.get(
        str(
            priority or ""
        ).lower(),
        4,
    )


def _append_key_value_table(
    document: Document,
    values: dict[str, Any],
) -> None:
    if not values:
        return

    table = document.add_table(
        rows=1,
        cols=2,
    )
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    header_cells[0].text = "Metric"
    header_cells[1].text = "Value"

    for key, value in values.items():
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = _value(value)


def _append_executive_summary(
    document: Document,
    *,
    executive_summary: dict[str, Any],
    financial_health: dict[str, Any],
    liquidity: dict[str, Any],
    funding_gap_summary: dict[str, Any],
) -> None:
    summary_financial_health = _mapping(
        executive_summary.get("financial_health")
    )
    summary_liquidity = _mapping(
        executive_summary.get("liquidity")
    )

    score = (
        summary_financial_health.get("score")
        if summary_financial_health.get("score") is not None
        else financial_health.get("score")
    )
    maximum = (
        summary_financial_health.get("maximum")
        if summary_financial_health.get("maximum") is not None
        else financial_health.get("maximum")
    )
    rating = (
        summary_financial_health.get("rating")
        or summary_financial_health.get("status")
        or financial_health.get("rating")
        or financial_health.get("status")
    )

    cash_runway_months = (
        summary_liquidity.get("cash_runway_months")
        if summary_liquidity.get("cash_runway_months") is not None
        else liquidity.get("cash_runway_months")
    )

    funding_gap_amount = funding_gap_summary.get(
        "funding_gap"
    )
    funding_coverage = funding_gap_summary.get(
        "applied_coverage_percentage"
    )

    table = document.add_table(
        rows=0,
        cols=3,
    )
    table.style = "Table Grid"

    rows = [
        (
            "Financial Health",
            (
                f"{_value(score)} / "
                f"{_value(maximum or 100)}"
                if score is not None
                else "-"
            ),
            _value(rating),
        ),
        (
            "Cash Runway",
            _value(cash_runway_months),
            "Months",
        ),
        (
            "Funding Gap",
            _currency(funding_gap_amount),
            _percentage(funding_coverage),
        ),
        (
            "High Current Risks",
            _value(
                executive_summary.get(
                    "high_current_risk_count"
                )
            ),
            "Current",
        ),
        (
            "High Forward Risks",
            _value(
                executive_summary.get(
                    "high_forward_risk_count"
                )
            ),
            "Forward-looking",
        ),
        (
            "High Opportunities",
            _value(
                executive_summary.get(
                    "high_opportunity_count"
                )
            ),
            "Opportunities",
        ),
        (
            "Priority Actions",
            _value(
                executive_summary.get(
                    "priority_action_count"
                )
            ),
            "Recommendations",
        ),
    ]

    for label, value, context in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = context


def _append_financial_health(
    document: Document,
    financial_health: dict[str, Any],
) -> None:
    score = financial_health.get("score")
    maximum = financial_health.get("maximum")
    rating = (
        financial_health.get("rating")
        or financial_health.get("status")
    )

    summary_table = document.add_table(
        rows=0,
        cols=3,
    )
    summary_table.style = "Table Grid"

    cells = summary_table.add_row().cells
    cells[0].text = "Overall Score"
    cells[1].text = (
        f"{_value(score)} / {_value(maximum or 100)}"
        if score is not None
        else "-"
    )
    cells[2].text = _value(rating)

    categories = _mapping(
        financial_health.get("categories")
    )

    if categories:
        document.add_heading(
            "Financial Health Areas",
            level=2,
        )

        category_table = document.add_table(
            rows=1,
            cols=4,
        )
        category_table.style = "Table Grid"

        header_cells = category_table.rows[0].cells
        header_cells[0].text = "Financial Health Area"
        header_cells[1].text = "Score"
        header_cells[2].text = "Maximum"
        header_cells[3].text = "Assessment"

        for key, item in categories.items():
            item_data = _mapping(item)

            cells = category_table.add_row().cells
            cells[0].text = _display_label(key)
            cells[1].text = _value(
                item_data.get("score")
            )
            cells[2].text = _value(
                item_data.get("maximum")
            )
            cells[3].text = _value(
                item_data.get("reason")
            )

    metrics = _mapping(
        financial_health.get("metrics")
    )

    selected_metrics = {
        key: metrics.get(key)
        for key in FINANCIAL_HEALTH_METRIC_KEYS
        if key in metrics
    }

    if selected_metrics:
        document.add_heading(
            "Key Financial Health Metrics",
            level=2,
        )
        _append_key_value_table(
            document,
            selected_metrics,
        )


def _append_liquidity(
    document: Document,
    liquidity: dict[str, Any],
) -> None:
    table = document.add_table(
        rows=0,
        cols=3,
    )
    table.style = "Table Grid"

    available_cash = liquidity.get(
        "available_cash"
    )
    blocked_cash = liquidity.get(
        "blocked_cash"
    )
    cash_runway_months = liquidity.get(
        "cash_runway_months"
    )

    rows = [
        (
            "Available Cash",
            _currency(available_cash),
            "Validated available liquidity",
        ),
        (
            "Blocked Cash",
            _currency(blocked_cash),
            "Restricted / blocked liquidity",
        ),
        (
            "Cash Runway",
            (
                f"{_value(cash_runway_months)} months"
                if cash_runway_months is not None
                else "-"
            ),
            "Operating-expense coverage",
        ),
    ]

    for label, value, context in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = context


def _append_budget(
    document: Document,
    budget: dict[str, Any],
) -> None:
    executive = _mapping(
        budget.get("executive_summary")
    )
    portfolio = _mapping(
        budget.get("portfolio_control")
    )

    source = executive or portfolio or budget

    cards = [
        (
            "Total Budget",
            _currency(
                source.get("total_budget")
            ),
            "Current validated portfolio",
        ),
        (
            "Total Actual",
            _currency(
                source.get("total_actual")
            ),
            "Validated actual expenditure",
        ),
    ]

    cards_table = document.add_table(
        rows=0,
        cols=3,
    )
    cards_table.style = "Table Grid"

    for label, value, context in cards:
        cells = cards_table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = context

    selected = {
        key: source.get(key)
        for key in BUDGET_SUMMARY_KEYS
        if key in source
    }

    if selected:
        document.add_heading(
            "Budget Metrics",
            level=2,
        )
        _append_key_value_table(
            document,
            selected,
        )

    organization = _mapping(
        budget.get("organization")
    )

    scope = {
        key: organization.get(key)
        for key in BUDGET_SCOPE_KEYS
        if key in organization
    }

    if scope:
        document.add_heading(
            "Organization Scope",
            level=2,
        )
        _append_key_value_table(
            document,
            scope,
        )


def _append_funding(
    document: Document,
    funding: dict[str, Any],
) -> None:
    funding_gap = _mapping(
        funding.get("funding_gap")
    )
    summary = _mapping(
        funding_gap.get("summary")
    )

    cards = [
        (
            "Remaining Requirement",
            _currency(
                summary.get(
                    "remaining_requirement"
                )
            ),
            "Validated uncovered requirement basis",
        ),
        (
            "Applied Secured Funding",
            _currency(
                summary.get(
                    "applied_secured_funding"
                )
            ),
            "Eligible secured funding applied",
        ),
        (
            "Funding Gap",
            _currency(
                summary.get(
                    "funding_gap"
                )
            ),
            (
                f"{_value(summary.get('applied_coverage_percentage'))}% coverage"
                if summary.get(
                    "applied_coverage_percentage"
                ) is not None
                else "Remaining uncovered need"
            ),
        ),
    ]

    cards_table = document.add_table(
        rows=0,
        cols=3,
    )
    cards_table.style = "Table Grid"

    for label, value, context in cards:
        cells = cards_table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = context

    selected = {
        key: summary.get(key)
        for key in FUNDING_SUMMARY_KEYS
        if key in summary
    }

    if selected:
        document.add_heading(
            "Funding Metrics",
            level=2,
        )
        _append_key_value_table(
            document,
            selected,
        )

    diagnostic_note = summary.get(
        "diagnostic_exposure_note"
    )

    if diagnostic_note:
        document.add_heading(
            "Funding Diagnostic Note",
            level=2,
        )
        document.add_paragraph(
            str(diagnostic_note)
        )

    diagnostics = _mapping(
        funding.get(
            "grant_diagnostics"
        )
    )

    diagnostic_counts: dict[str, Any] = {}

    for key in GRANT_DIAGNOSTIC_KEYS:
        value = diagnostics.get(key)

        if isinstance(value, list):
            diagnostic_counts[key] = len(value)
        elif isinstance(value, int):
            diagnostic_counts[key] = value

    if diagnostic_counts:
        document.add_heading(
            "Grant Diagnostics",
            level=2,
        )
        _append_key_value_table(
            document,
            diagnostic_counts,
        )


def _append_core_cost_coverage(
    document: Document,
    core_cost_coverage: dict[str, Any],
) -> None:
    summary = _mapping(
        core_cost_coverage.get("summary")
    )

    coverage_cards = [
        (
            "Needed Core Cost",
            _currency(
                summary.get(
                    "needed_core_cost"
                )
            ),
            "Validated core cost requirement",
        ),
        (
            "Direct Grant Coverage",
            _currency(
                summary.get(
                    "direct_grant_coverage"
                )
            ),
            "Core costs covered directly by grants",
        ),
        (
            "Allocated Indirect Recovery",
            _currency(
                summary.get(
                    "allocated_indirect_recovery"
                )
            ),
            "Indirect recovery allocated by management",
        ),
        (
            "Unrestricted / Core Funding",
            _currency(
                summary.get(
                    "unrestricted_core_funding"
                )
            ),
            "Other unrestricted or core funding",
        ),
    ]

    coverage_table = document.add_table(
        rows=0,
        cols=3,
    )
    coverage_table.style = "Table Grid"

    for label, value, context in coverage_cards:
        cells = coverage_table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = context

    status_cards = [
        (
            "Remaining Core Cost Gap",
            _currency(
                summary.get(
                    "remaining_core_cost_gap"
                )
            ),
            "Remaining uncovered core cost requirement",
        ),
        (
            "Core Cost Coverage",
            _percentage(
                summary.get(
                    "core_cost_coverage_percentage"
                )
            ),
            "Verified percentage of core costs covered",
        ),
        (
            "Available Indirect Recovery",
            _currency(
                summary.get(
                    "available_indirect_recovery"
                )
            ),
            "Recovered indirect funding available",
        ),
        (
            "Used / Charged Indirect",
            _currency(
                summary.get(
                    "used_indirect_recovery"
                )
            ),
            "Indirect recovery actually used or charged",
        ),
    ]

    status_table = document.add_table(
        rows=0,
        cols=3,
    )
    status_table.style = "Table Grid"

    for label, value, context in status_cards:
        cells = status_table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        cells[2].text = context

    lines = core_cost_coverage.get("lines")

    if isinstance(lines, list) and lines:
        document.add_heading(
            "Core Cost Coverage Lines",
            level=2,
        )

        line_table = document.add_table(
            rows=1,
            cols=6,
        )
        line_table.style = "Table Grid"

        headers = [
            "Budget Line",
            "Needed",
            "Direct",
            "Allocated Indirect",
            "Unrestricted",
            "Remaining Gap",
        ]

        for index, header in enumerate(headers):
            line_table.rows[0].cells[index].text = header

        for item in lines:
            line = _mapping(item)

            budget_line_code = _value(
                line.get("budget_line_code")
            )
            budget_line_name = _value(
                line.get("budget_line_name")
            )

            if (
                budget_line_code != "-"
                and budget_line_name != "-"
            ):
                budget_line = (
                    f"{budget_line_code} - "
                    f"{budget_line_name}"
                )
            elif budget_line_code != "-":
                budget_line = budget_line_code
            else:
                budget_line = budget_line_name

            cells = line_table.add_row().cells
            cells[0].text = budget_line
            cells[1].text = _currency(
                line.get("needed_core_cost")
            )
            cells[2].text = _currency(
                line.get("direct_grant_coverage")
            )
            cells[3].text = _currency(
                line.get(
                    "allocated_indirect_recovery"
                )
            )
            cells[4].text = _currency(
                line.get(
                    "unrestricted_core_funding"
                )
            )
            cells[5].text = _currency(
                line.get(
                    "remaining_core_cost_gap"
                )
            )


def _append_risk_cards(
    document: Document,
    risks: list[dict[str, Any]],
) -> None:
    """
    Render validated risk items without changing or
    re-evaluating their severity or evidence.
    """
    for risk in risks:
        title = (
            risk.get("title")
            or "Financial Risk"
        )

        document.add_heading(
            str(title),
            level=3,
        )

        table = document.add_table(
            rows=0,
            cols=2,
        )
        table.style = "Table Grid"

        fields = [
            ("Severity", risk.get("severity")),
            ("Category", risk.get("category")),
            ("Evidence", risk.get("evidence")),
            (
                "Recommendation",
                risk.get("recommendation"),
            ),
            ("Action", risk.get("action")),
        ]

        for label, value in fields:
            if value is None or value == "":
                continue

            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = _value(value)


def _append_risks(
    document: Document,
    risks: dict[str, Any],
) -> None:
    """
    Render validated current and forward-looking risks.

    This function is presentation-only and does not score,
    reclassify, or generate risks.
    """
    current = _list_of_dicts(
        risks.get("current")
    )
    forward = _list_of_dicts(
        risks.get("forward")
    )

    document.add_heading(
        "Current Financial Risks",
        level=2,
    )

    if current:
        _append_risk_cards(
            document,
            current,
        )
    else:
        document.add_paragraph(
            "No current financial risks available."
        )

    document.add_heading(
        "Forward-Looking Risks",
        level=2,
    )

    if forward:
        _append_risk_cards(
            document,
            forward,
        )
    else:
        document.add_paragraph(
            "No forward-looking risks available."
        )

def _append_opportunities(
    document: Document,
    opportunities: dict[str, Any],
) -> None:
    """
    Render validated financial opportunities.

    This function is presentation-only and does not generate,
    rank, or reclassify opportunities.
    """
    items = _list_of_dicts(
        opportunities.get("all")
    )

    document.add_heading(
        "Financial Opportunities",
        level=2,
    )

    if not items:
        document.add_paragraph(
            "No financial opportunities available."
        )
        return

    for opportunity in items:
        title = (
            opportunity.get("title")
            or "Financial Opportunity"
        )

        document.add_heading(
            str(title),
            level=3,
        )

        table = document.add_table(
            rows=0,
            cols=2,
        )
        table.style = "Table Grid"

        fields = [
            (
                "Priority",
                opportunity.get("priority"),
            ),
            (
                "Category",
                opportunity.get("category"),
            ),
            (
                "Evidence",
                opportunity.get("evidence"),
            ),
            (
                "Recommended Action",
                opportunity.get(
                    "recommended_action"
                ),
            ),
            (
                "Action",
                opportunity.get("action"),
            ),
        ]

        for label, value in fields:
            if value is None or value == "":
                continue

            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = _value(value)

def _append_recommendations(
    document: Document,
    recommendations: dict[str, Any],
) -> None:
    """
    Render validated CFO recommendations.

    This function is presentation-only and does not generate,
    reprioritize, or modify recommendations.
    """
    items = _list_of_dicts(
        recommendations.get("all")
    )

    if not items:
        document.add_paragraph(
            "No CFO recommendations available."
        )
        return

    ordered = sorted(
        items,
        key=lambda item: _priority_rank(
            item.get("priority")
        ),
    )

    for number, recommendation in enumerate(
        ordered,
        start=1,
    ):
        title = (
            recommendation.get("title")
            or recommendation.get("action")
            or f"Management Action {number}"
        )

        document.add_heading(
            str(title),
            level=3,
        )

        table = document.add_table(
            rows=0,
            cols=2,
        )
        table.style = "Table Grid"

        fields = [
            (
                "Priority",
                recommendation.get("priority"),
            ),
            (
                "Category",
                recommendation.get("category"),
            ),
            (
                "Evidence",
                recommendation.get("evidence"),
            ),
            (
                "Reason",
                recommendation.get("reason"),
            ),
            (
                "Expected Impact",
                recommendation.get(
                    "expected_impact"
                ),
            ),
            (
                "Action",
                recommendation.get("action"),
            ),
        ]

        for label, value in fields:
            if value is None or value == "":
                continue

            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = _value(value)

def _append_period_comparison(
    document: Document,
    comparison: dict[str, Any],
) -> None:
    """
    Render a validated financial period comparison.

    This function is presentation-only and does not calculate
    trend values or change percentages.
    """
    current_period = comparison.get(
        "current_period"
    )
    previous_period = comparison.get(
        "previous_period"
    )

    period_table = document.add_table(
        rows=0,
        cols=2,
    )
    period_table.style = "Table Grid"

    cells = period_table.add_row().cells
    cells[0].text = "Current"
    cells[1].text = _value(current_period)

    cells = period_table.add_row().cells
    cells[0].text = "Previous"
    cells[1].text = _value(previous_period)

    rows: list[
        tuple[str, dict[str, Any]]
    ] = []

    for key in (
        "revenue",
        "expenses",
        "net_result",
    ):
        item = _mapping(
            comparison.get(key)
        )

        if item:
            rows.append(
                (
                    key,
                    item,
                )
            )

    if not rows:
        return

    table = document.add_table(
        rows=1,
        cols=5,
    )
    table.style = "Table Grid"

    headers = [
        "Metric",
        "Current",
        "Previous",
        "Change %",
        "Direction",
    ]

    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header

    for key, item in rows:
        cells = table.add_row().cells

        cells[0].text = _display_label(key)
        cells[1].text = _currency(
            item.get("current_value")
        )
        cells[2].text = _currency(
            item.get("previous_value")
        )
        cells[3].text = _percentage(
            item.get("change_percentage")
        )
        cells[4].text = _value(
            item.get("direction")
        )


def _append_trends(
    document: Document,
    trends: dict[str, Any],
) -> None:
    """
    Render validated Financial Trend Intelligence.

    This function consumes existing trend intelligence and does
    not recalculate financial trends or period comparisons.
    """
    coverage = _mapping(
        trends.get("coverage")
    )
    quality = _mapping(
        trends.get("quality")
    )
    latest_month = _mapping(
        trends.get(
            "latest_month_comparison"
        )
    )
    latest_year = _mapping(
        trends.get(
            "latest_year_comparison"
        )
    )

    if coverage:
        document.add_heading(
            "Trend Coverage",
            level=2,
        )
        _append_key_value_table(
            document,
            coverage,
        )

    if quality:
        document.add_heading(
            "Data Quality",
            level=2,
        )
        _append_key_value_table(
            document,
            quality,
        )

    if latest_month:
        document.add_heading(
            "Latest Month Comparison",
            level=2,
        )
        _append_period_comparison(
            document,
            latest_month,
        )

    if latest_year:
        document.add_heading(
            "Latest Year Comparison",
            level=2,
        )
        _append_period_comparison(
            document,
            latest_year,
        )

def _append_forecast(
    document: Document,
    forecast: dict[str, Any],
) -> None:
    """
    Render validated Financial Forecast Intelligence.

    This function is presentation-only and does not calculate
    or modify forecast values.
    """
    baseline = _mapping(
        forecast.get("baseline")
    )
    totals = _mapping(
        forecast.get("forecast_totals")
    )
    confidence = _mapping(
        forecast.get("confidence")
    )

    if baseline:
        document.add_heading(
            "Baseline",
            level=2,
        )
        _append_key_value_table(
            document,
            baseline,
        )

    if totals:
        cards = [
            (
                "Forecast Revenue",
                _currency(
                    totals.get("revenue")
                ),
                "Forecast horizon total",
            ),
            (
                "Forecast Expenses",
                _currency(
                    totals.get("expenses")
                ),
                "Forecast horizon total",
            ),
            (
                "Forecast Net Result",
                _currency(
                    totals.get("net_result")
                ),
                "Forecast horizon result",
            ),
        ]

        table = document.add_table(
            rows=0,
            cols=3,
        )
        table.style = "Table Grid"

        for label, value, context in cards:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
            cells[2].text = context

    if confidence:
        document.add_heading(
            "Forecast Confidence",
            level=2,
        )
        _append_key_value_table(
            document,
            confidence,
        )

def _append_methodology(
    document: Document,
    methodology: dict[str, Any],
    funding: dict[str, Any],
) -> None:
    """
    Render report methodology and evidence principles.

    This function is presentation-only. It documents the methodology
    already supplied by validated AI-FOS report outputs and does not
    recalculate financial results.
    """
    if methodology:
        _append_key_value_table(
            document,
            methodology,
        )

    funding_gap = _mapping(
        funding.get("funding_gap")
    )

    for possible_key in (
        "methodology",
        "metadata",
    ):
        funding_methodology = _mapping(
            funding_gap.get(
                possible_key
            )
        )

        if funding_methodology:
            document.add_heading(
                "Funding Gap Methodology",
                level=2,
            )

            _append_key_value_table(
                document,
                funding_methodology,
            )

            break

    document.add_heading(
        "Evidence Principle",
        level=2,
    )

    document.add_paragraph(
        "AI-FOS report presentation consumes validated financial "
        "model outputs. The Word renderer does not independently "
        "recalculate financial results."
    )

def generate_cfo_report_word(
    report: dict[str, Any],
) -> bytes:
    """
    Generate the AI-FOS CFO Financial Intelligence Word report.

    The renderer is presentation-only.

    It consumes the validated structured CFO report and does not
    recalculate financial results.
    """
    document = Document()

    executive_summary = _mapping(
        report.get("executive_summary")
    )
    financial_health = _mapping(
        report.get("financial_health")
    )
    liquidity = _mapping(
        report.get("liquidity")
    )
    budget = _mapping(
        report.get("budget")
    )
    funding = _mapping(
        report.get("funding")
    )
    core_cost_coverage = _mapping(
        report.get("core_cost_coverage")
    )
    risks = _mapping(
        report.get("risks")
    )

    opportunities = _mapping(
        report.get("opportunities")
    )

    recommendations = _mapping(
        report.get("recommendations")
    ) 

    trends = _mapping(
        report.get("trends")
    )  

    forecast = _mapping(report.get("forecast"))  

    methodology = _mapping(report.get("methodology"))               

    funding_gap = _mapping(
        funding.get("funding_gap")
    )
    funding_gap_summary = _mapping(
        funding_gap.get("summary")
    )

    document.add_heading(
        "CFO Financial Intelligence Report",
        level=0,
    )

    section_titles = [
        "Executive Summary",
        "Financial Health",
        "Liquidity",
        "Budget Performance",
        "Funding & Grants",
        "Core Cost Coverage",
        "Current & Forward Risks",
        "Opportunities",
        "CFO Recommendations",
        "Trends",
        "Forecast",
        "Methodology",
    ]

    for section_title in section_titles:
        document.add_heading(
            section_title,
            level=1,
        )

        if section_title == "Executive Summary":
            _append_executive_summary(
                document,
                executive_summary=executive_summary,
                financial_health=financial_health,
                liquidity=liquidity,
                funding_gap_summary=funding_gap_summary,
            )

        elif section_title == "Financial Health":
            _append_financial_health(
                document,
                financial_health,
            )

        elif section_title == "Liquidity":
            _append_liquidity(
                document,
                liquidity,
            )

        elif section_title == "Budget Performance":
            _append_budget(
                document,
                budget,
            )

        elif section_title == "Funding & Grants":
            _append_funding(
                document,
                funding,
            )

        elif section_title == "Core Cost Coverage":
            _append_core_cost_coverage(
                document,
                core_cost_coverage,
            )

        elif section_title == "Current & Forward Risks":
            _append_risks(
                document,
                risks,
            )

        elif section_title == "Opportunities":
            _append_opportunities(
                document,
                opportunities,
            ) 

        elif section_title == "CFO Recommendations":
            _append_recommendations(
                document,
                recommendations,
            )

        elif section_title == "Trends":
            _append_trends(
                document,
                trends,
            )

        elif section_title == "Forecast":
            _append_forecast(
                document,
                forecast,
            )

        elif section_title == "Methodology":
            _append_methodology(
                document,
                methodology,
                funding,
            )                                                           

    buffer = BytesIO()
    document.save(buffer)

    word_bytes = buffer.getvalue()
    buffer.close()

    return word_bytes