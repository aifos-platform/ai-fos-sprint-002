from io import BytesIO

from openpyxl import load_workbook

from app.services.standard_income_statement_excel import (
    generate_standard_income_statement_excel,
)


def _sample_report() -> dict:
    return {
        "status": "available",
        "period": {
            "start_date": "2026-01-01",
            "end_date": "2026-12-31",
        },
        "sections": [
            {
                "name": "Revenue",
                "lines": [
                    {
                        "account_number": "700000",
                        "account_name": "Revenue",
                        "level": 1,
                        "amount": 1000.0,
                        "is_posting_account": False,
                    },
                    {
                        "account_number": "701000",
                        "account_name": "Grant Revenue",
                        "level": 2,
                        "amount": 1000.0,
                        "is_posting_account": True,
                    },
                ],
                "total": 1000.0,
                "total_label": "Total Revenue",
            },
            {
                "name": "Expenses",
                "lines": [
                    {
                        "account_number": "600000",
                        "account_name": "Expenses",
                        "level": 1,
                        "amount": 500.0,
                        "is_posting_account": False,
                    },
                    {
                        "account_number": "601000",
                        "account_name": "Personnel",
                        "level": 2,
                        "amount": 500.0,
                        "is_posting_account": True,
                    },
                ],
                "total": 500.0,
                "total_label": "Total Expenses",
            },
        ],
        "totals": {
            "revenue": 1000.0,
            "expenses": 500.0,
            "net_result": 500.0,
        },
        "reconciliation": {
            "reconciled": True,
        },
        "controls": {
            "reporting_layer_only": True,
            "financial_recalculation_performed": False,
        },
    }


def _load_generated_workbook():
    excel_bytes = generate_standard_income_statement_excel(
        _sample_report(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    return load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )


def _worksheet_values(worksheet) -> list:
    return [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]


def test_income_statement_excel_returns_valid_workbook():
    excel_bytes = generate_standard_income_statement_excel(
        _sample_report(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    workbook = load_workbook(
        BytesIO(excel_bytes)
    )

    assert workbook.sheetnames == [
        "Income Statement"
    ]


def test_income_statement_excel_contains_report_identity():
    workbook = _load_generated_workbook()

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    assert "Test NGO" in values
    assert "Income Statement" in values

    assert (
        "For the period "
        "2026-01-01 to 2026-12-31"
    ) in values

    assert "Currency: USD" in values


def test_income_statement_excel_preserves_report_lines():
    workbook = _load_generated_workbook()

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    assert "700000" in values
    assert "Revenue" in values

    assert "701000" in values
    assert "Grant Revenue" in values

    assert "600000" in values
    assert "Expenses" in values

    assert "601000" in values
    assert "Personnel" in values


def test_income_statement_excel_preserves_report_amounts():
    workbook = _load_generated_workbook()

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    numeric_values = [
        value
        for value in values
        if isinstance(
            value,
            (int, float),
        )
    ]

    assert 1000.0 in numeric_values
    assert 500.0 in numeric_values


def test_income_statement_excel_contains_section_totals():
    workbook = _load_generated_workbook()

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    assert "Total Revenue" in values
    assert "Total Expenses" in values


def test_income_statement_excel_contains_net_result():
    workbook = _load_generated_workbook()

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    assert (
        "NET SURPLUS / (DEFICIT)"
        in values
    )

    net_result_rows = [
        row
        for row in worksheet.iter_rows()
        if row[0].value
        == "NET SURPLUS / (DEFICIT)"
    ]

    assert len(net_result_rows) == 1

    assert (
        net_result_rows[0][3].value
        == 500.0
    )


def test_income_statement_excel_contains_reconciliation_note():
    workbook = _load_generated_workbook()

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    assert (
        "AI-FOS Standard Financial Report"
        " • Reconciled to validated Income Statement"
    ) in values


def test_income_statement_excel_does_not_mutate_source_report():
    report = _sample_report()

    original_report = {
        "status": report["status"],
        "period": dict(
            report["period"]
        ),
        "sections": [
            {
                **section,
                "lines": [
                    dict(line)
                    for line in section["lines"]
                ],
            }
            for section in report["sections"]
        ],
        "totals": dict(
            report["totals"]
        ),
        "reconciliation": dict(
            report["reconciliation"]
        ),
        "controls": dict(
            report["controls"]
        ),
    }

    generate_standard_income_statement_excel(
        report,
        organisation_name="Test NGO",
        base_currency="USD",
    )

    assert report == original_report


def test_income_statement_excel_is_presentation_only():
    report = _sample_report()

    excel_bytes = generate_standard_income_statement_excel(
        report,
        organisation_name="Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    worksheet = workbook[
        "Income Statement"
    ]

    formulas = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if (
            isinstance(cell.value, str)
            and cell.value.startswith("=")
        )
    ]

    assert formulas == []


def test_income_statement_excel_handles_empty_report():
    excel_bytes = generate_standard_income_statement_excel(
        {},
        organisation_name="Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes)
    )

    assert workbook.sheetnames == [
        "Income Statement"
    ]

    worksheet = workbook[
        "Income Statement"
    ]

    values = _worksheet_values(
        worksheet
    )

    assert "Test NGO" in values
    assert "Income Statement" in values