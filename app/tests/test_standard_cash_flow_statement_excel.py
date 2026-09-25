from io import BytesIO

from openpyxl import load_workbook

from app.services.standard_cash_flow_statement_excel import (
    generate_standard_cash_flow_statement_excel,
)


def _standard_cash_flow_statement() -> dict:
    return {
        "statement_type": "standard_cash_flow_statement",
        "version": "v1",
        "method": "transaction_counter_account_analysis",
        "sections": {
            "operating_activities": {
                "label": (
                    "Cash Flows from Operating Activities"
                ),
                "amount": 1000.0,
            },
            "investing_activities": {
                "label": (
                    "Cash Flows from Investing Activities"
                ),
                "amount": -250.0,
            },
            "financing_activities": {
                "label": (
                    "Cash Flows from Financing Activities"
                ),
                "amount": 100.0,
            },
        },
        "net_change_in_cash": 850.0,
        "unclassified_cash_movement": 0.0,
        "reconciliation": {
            "classified_net_change": 850.0,
            "net_change_in_cash": 850.0,
            "difference": 0.0,
            "is_reconciled": True,
        },
        "cash_accounts": [],
        "classification_diagnostics": {
            "document_count": 20,
            "classified_document_count": 20,
            "unclassified_document_count": 0,
            "activity_detail_count": 20,
        },
        "validation": {
            "status": "validated",
            "reconciled": True,
            "has_unclassified_cash_movement": False,
        },
        "note": (
            "Standard Cash Flow Statement v1 uses the "
            "verified AI-FOS transaction-level cash-flow "
            "classification as its source."
        ),
    }


def test_standard_cash_flow_statement_excel_generates_workbook():
    report = _standard_cash_flow_statement()

    excel_bytes = (
        generate_standard_cash_flow_statement_excel(
            report=report,
            organisation_name="Test NGO",
            base_currency="USD",
        )
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    assert workbook.sheetnames == [
        "Cash Flow Statement"
    ]

    worksheet = workbook[
        "Cash Flow Statement"
    ]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "Test NGO" in values
    assert "Cash Flow Statement" in values
    assert "Currency: USD" in values

    assert (
        "Cash Flows from Operating Activities"
        in values
    )

    assert (
        "Cash Flows from Investing Activities"
        in values
    )

    assert (
        "Cash Flows from Financing Activities"
        in values
    )

    assert "Unclassified Cash Movement" in values
    assert "NET CHANGE IN CASH" in values
    assert "RECONCILIATION" in values
    assert "CLASSIFICATION DIAGNOSTICS" in values

    assert 1000.0 in values
    assert -250.0 in values
    assert 100.0 in values
    assert 850.0 in values


def test_standard_cash_flow_statement_excel_preserves_report_values():
    report = _standard_cash_flow_statement()

    report["sections"][
        "operating_activities"
    ]["amount"] = 1234.56

    report["unclassified_cash_movement"] = 45.67
    report["net_change_in_cash"] = 1130.23

    excel_bytes = (
        generate_standard_cash_flow_statement_excel(
            report=report,
        )
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    worksheet = workbook[
        "Cash Flow Statement"
    ]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert 1234.56 in values
    assert 45.67 in values
    assert 1130.23 in values


def test_standard_cash_flow_statement_excel_does_not_add_unverified_balances():
    report = _standard_cash_flow_statement()

    excel_bytes = (
        generate_standard_cash_flow_statement_excel(
            report=report,
        )
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    worksheet = workbook[
        "Cash Flow Statement"
    ]

    text_values = [
        str(cell.value)
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert not any(
        "Opening Cash Balance" in value
        for value in text_values
    )

    assert not any(
        "Closing Cash Balance" in value
        for value in text_values
    )