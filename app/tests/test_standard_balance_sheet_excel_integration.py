from io import BytesIO

from openpyxl import load_workbook

from app.services.cfo_report_excel import (
    generate_cfo_report_excel,
)


def _standard_balance_sheet():
    return {
        "status": "available",
        "report_type": "balance_sheet",
        "title": "Balance Sheet",
        "reporting_basis": "as_of_date",
        "sections": [
            {
                "section": "Assets",
                "lines": [
                    {
                        "account_number": "201100",
                        "account_name": "Cash",
                        "parent_account": "201000",
                        "level": 2,
                        "financial_category": "Asset",
                        "financial_subcategory": "Cash",
                        "normal_balance": "Debit",
                        "amount": 1500.0,
                    },
                ],
                "total": 1500.0,
            },
            {
                "section": "Liabilities",
                "lines": [
                    {
                        "account_number": "401100",
                        "account_name": "Accounts Payable",
                        "parent_account": "401000",
                        "level": 2,
                        "financial_category": "Liability",
                        "financial_subcategory": "Payables",
                        "normal_balance": "Credit",
                        "amount": 900.0,
                    },
                ],
                "total": 900.0,
            },
            {
                "section": "Equity / Net Assets",
                "lines": [
                    {
                        "account_number": "101100",
                        "account_name": "Equity",
                        "parent_account": "101000",
                        "level": 2,
                        "financial_category": "Equity",
                        "financial_subcategory": "Equity",
                        "normal_balance": "Credit",
                        "amount": 400.0,
                    },
                ],
                "reported_equity": 400.0,
                "current_period_result": 200.0,
                "total": 600.0,
            },
        ],
        "totals": {
            "assets": 1500.0,
            "liabilities": 900.0,
            "reported_equity": 400.0,
            "current_period_result": 200.0,
            "equity": 600.0,
            "liabilities_and_equity": 1500.0,
            "difference": 0.0,
        },
        "validated_totals": {
            "assets": 1500.0,
            "liabilities": 900.0,
            "reported_equity": 400.0,
            "current_period_result": 200.0,
            "equity": 600.0,
            "liabilities_and_equity": 1500.0,
            "difference": 0.0,
        },
        "reconciliation": {
            "reconciled": True,
        },
        "controls": {
            "reporting_layer_only": True,
            "validated_balance_sheet_preserved": True,
            "financial_recalculation_performed": False,
            "missing_accounts_not_invented": True,
            "reconciliation_required": True,
        },
    }


def _minimal_cfo_report():
    return {
        "executive_summary": {},
        "financial_health": {},
        "liquidity": {},
        "budget": {},
        "funding": {},
        "risks": {},
        "opportunities": {},
        "recommendations": {},
        "trends": {},
        "forecast": {},
        "methodology": {},
    }


def test_cfo_excel_includes_standard_balance_sheet():
    excel_bytes = generate_cfo_report_excel(
        _minimal_cfo_report(),
        standard_balance_sheet=_standard_balance_sheet(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    assert "Balance Sheet" in workbook.sheetnames

    worksheet = workbook["Balance Sheet"]

    assert worksheet["A1"].value == "Test NGO"
    assert worksheet["A2"].value == "Balance Sheet"


def test_cfo_excel_balance_sheet_preserves_values():
    excel_bytes = generate_cfo_report_excel(
        _minimal_cfo_report(),
        standard_balance_sheet=_standard_balance_sheet(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    worksheet = workbook["Balance Sheet"]

    values = [
        [
            cell.value
            for cell in row
        ]
        for row in worksheet.iter_rows()
    ]

    rows_by_label = {
        row[0]: row
        for row in values
        if row and row[0] is not None
    }

    assert rows_by_label["Total Assets"][3] == 1500.0
    assert rows_by_label["Total Liabilities"][3] == 900.0

    assert (
        rows_by_label["Adjusted Equity / Net Assets"][3]
        == 600.0
    )

    assert (
        rows_by_label["TOTAL LIABILITIES & EQUITY"][3]
        == 1500.0
    )


def test_cfo_excel_balance_sheet_is_presentation_only():
    excel_bytes = generate_cfo_report_excel(
        _minimal_cfo_report(),
        standard_balance_sheet=_standard_balance_sheet(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    worksheet = workbook["Balance Sheet"]

    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith("=")


def test_cfo_excel_works_without_balance_sheet():
    excel_bytes = generate_cfo_report_excel(
        _minimal_cfo_report(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )

    assert "Balance Sheet" not in workbook.sheetnames