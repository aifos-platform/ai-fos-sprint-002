from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO

from app.main import app
from app.organization import Organization
from app.services.financial_model_service import (
    FinancialModelService,
)


client = TestClient(app)


def _workspace(
    financial_model_folder: Path,
) -> dict:
    return {
        "paths": {
            "financial_model": str(
                financial_model_folder
            ),
        },
    }


def _standard_cash_flow_statement() -> dict:
    return {
        "status": "available",
        "report_type": "standard_cash_flow_statement",
        "totals": {
            "operating_activities": 1000.0,
            "investing_activities": -300.0,
            "financing_activities": 200.0,
            "net_change_in_cash": 900.0,
        },
        "sections": {
            "operating_activities": {
                "amount": 1000.0,
                "lines": [],
            },
            "investing_activities": {
                "amount": -300.0,
                "lines": [],
            },
            "financing_activities": {
                "amount": 200.0,
                "lines": [],
            },
        },
        "reconciliation": {
            "reconciled": True,
        },
        "controls": {
            "reporting_layer_only": True,
            "financial_recalculation_performed": False,
        },
    }


def _organization_metadata() -> dict:
    return {
        "id": "test-org",
        "name": "Test NGO",
        "base_currency": "EUR",
        "active": True,
    }


def test_standard_cash_flow_statement_excel_api_returns_workbook(
    monkeypatch,
    tmp_path: Path,
):
    report = _standard_cash_flow_statement()
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organization_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_cash_flow_statement.json",
        report,
    )

    response = client.get(
        "/reports/test-org/cash-flow/excel"
    )

    assert response.status_code == 200

    assert (
        response.headers["content-type"]
        == (
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    assert (
        'filename="AI-FOS_test-org_Cash_Flow_Statement.xlsx"'
        in response.headers["content-disposition"]
    )

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    assert workbook.sheetnames == [
        "Cash Flow Statement"
    ]

    worksheet = workbook[
        "Cash Flow Statement"
    ]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "Test NGO" in values
    assert "Cash Flow Statement" in values


def test_standard_cash_flow_statement_excel_api_uses_registry_metadata(
    monkeypatch,
    tmp_path: Path,
):
    report = _standard_cash_flow_statement()
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: {
            "id": "test-org",
            "name": "Global Relief NGO",
            "base_currency": "GBP",
            "active": True,
        },
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_cash_flow_statement.json",
        report,
    )

    response = client.get(
        "/reports/test-org/cash-flow/excel"
    )

    assert response.status_code == 200

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    worksheet = workbook[
        "Cash Flow Statement"
    ]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "Global Relief NGO" in values

    currency_values = [
        str(value)
        for value in values
        if value is not None
    ]

    assert any(
        "GBP" in value
        for value in currency_values
    )


def test_standard_cash_flow_statement_excel_api_missing_workspace_returns_404(
    monkeypatch,
):
    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: None,
    )

    response = client.get(
        "/reports/missing-org/cash-flow/excel"
    )

    assert response.status_code == 404

    assert (
        response.json()["detail"]
        == "No workspace found for this organization."
    )


def test_standard_cash_flow_statement_excel_api_missing_report_returns_404(
    monkeypatch,
    tmp_path: Path,
):
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    response = client.get(
        "/reports/test-org/cash-flow/excel"
    )

    assert response.status_code == 404

    assert (
        "Standard Cash Flow Statement is not available yet"
        in response.json()["detail"]
    )


def test_standard_cash_flow_statement_excel_api_does_not_recalculate_financials(
    monkeypatch,
    tmp_path: Path,
):
    report = _standard_cash_flow_statement()
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organization_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_cash_flow_statement.json",
        report,
    )

    def forbidden_recalculation(
        *args,
        **kwargs,
    ):
        raise AssertionError(
            "Cash Flow Statement Excel API must not "
            "recalculate financials."
        )

    monkeypatch.setattr(
        Organization,
        "process_financials",
        forbidden_recalculation,
    )

    response = client.get(
        "/reports/test-org/cash-flow/excel"
    )

    assert response.status_code == 200

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    assert workbook.sheetnames == [
        "Cash Flow Statement"
    ]