from openpyxl import Workbook

from app.services.excel_reader import (
    inspect_workbook,
)


def test_budget_workbook_reads_core_cost_coverage_sheet(
    tmp_path,
):
    file_path = (
        tmp_path
        / "budget_with_core_cost_coverage.xlsx"
    )

    workbook = Workbook()

    #
    # First sheet must be detectable as a Budget
    # workbook by AI-FOS.
    #
    budget_sheet = workbook.active
    budget_sheet.title = "Available Budget"

    budget_sheet.append(
        [
            "Budget Line Code",
            "Budget Line Code Name",
            "Budget Notes",
            "Total Original Budget (USD)",
        ]
    )

    budget_sheet.append(
        [
            "SAL-001",
            "Salaries",
            "Core personnel cost",
            100000.0,
        ]
    )

    #
    # Optional Core Cost Coverage worksheet.
    #
    coverage_sheet = workbook.create_sheet(
        "Core Cost Coverage"
    )

    coverage_sheet.append(
        [
            "Coverage Type",
            "Fund Code",
            "Budget Line Code",
            "Amount",
        ]
    )

    coverage_sheet.append(
        [
            "direct_grant_coverage",
            "GRANT-001",
            "SAL-001",
            25000.0,
        ]
    )

    coverage_sheet.append(
        [
            None,
            None,
            None,
            None,
        ]
    )    

    workbook.save(file_path)

    result = inspect_workbook(
        str(file_path)
    )

    assert (
        result[
            "core_cost_coverage_sheet_name"
        ]
        == "Core Cost Coverage"
    )

    assert len(
        result[
            "core_cost_coverage_lines"
        ]
    ) == 1

    coverage = result[
        "core_cost_coverage_lines"
    ][0]

    assert (
        coverage["coverage_type"]
        == "direct_grant_coverage"
    )

    assert (
        coverage["fund_code"]
        == "GRANT-001"
    )

    assert (
        coverage["budget_line_code"]
        == "SAL-001"
    )

    assert coverage["amount"] == 25000.0

    assert coverage["requires_review"] is False

def test_budget_workbook_without_core_cost_coverage_sheet_returns_empty_data(
    tmp_path,
):
    file_path = (
        tmp_path
        / "budget_without_core_cost_coverage.xlsx"
    )

    workbook = Workbook()

    budget_sheet = workbook.active
    budget_sheet.title = "Available Budget"

    budget_sheet.append(
        [
            "Budget Line Code",
            "Budget Line Code Name",
            "Budget Notes",
            "Total Original Budget (USD)",
        ]
    )

    budget_sheet.append(
        [
            "SAL-001",
            "Salaries",
            "Core personnel cost",
            100000.0,
        ]
    )

    workbook.save(file_path)

    result = inspect_workbook(
        str(file_path)
    )

    assert (
        result[
            "core_cost_coverage_sheet_name"
        ]
        is None
    )

    assert (
        result[
            "core_cost_coverage_mapping"
        ]
        == {}
    )

    assert (
        result[
            "core_cost_coverage_lines"
        ]
        == []
    )

def test_budget_workbook_reads_core_cost_coverage_alias_headers(
    tmp_path,
):
    file_path = (
        tmp_path
        / "budget_with_core_cost_coverage_aliases.xlsx"
    )

    workbook = Workbook()

    budget_sheet = workbook.active
    budget_sheet.title = "Available Budget"

    budget_sheet.append(
        [
            "Budget Line Code",
            "Budget Line Code Name",
            "Budget Notes",
            "Total Original Budget (USD)",
        ]
    )

    budget_sheet.append(
        [
            "SAL-001",
            "Salaries",
            "Core personnel cost",
            100000.0,
        ]
    )

    coverage_sheet = workbook.create_sheet(
        "Core Cost Coverage"
    )

    coverage_sheet.append(
        [
            "Core Cost Coverage Type",
            "Grant Code",
            "Budget Line Code",
            "Coverage Amount",
        ]
    )

    coverage_sheet.append(
        [
            "direct_grant_coverage",
            "GRANT-001",
            "SAL-001",
            30000.0,
        ]
    )

    workbook.save(file_path)

    result = inspect_workbook(
        str(file_path)
    )

    assert (
        result[
            "core_cost_coverage_sheet_name"
        ]
        == "Core Cost Coverage"
    )

    assert result[
        "core_cost_coverage_mapping"
    ] == {
        "coverage_type": "Core Cost Coverage Type",
        "fund_code": "Grant Code",
        "budget_line_code": "Budget Line Code",
        "amount": "Coverage Amount",
    }

    assert len(
        result[
            "core_cost_coverage_lines"
        ]
    ) == 1

    coverage = result[
        "core_cost_coverage_lines"
    ][0]

    assert (
        coverage["coverage_type"]
        == "direct_grant_coverage"
    )

    assert (
        coverage["fund_code"]
        == "GRANT-001"
    )

    assert (
        coverage["budget_line_code"]
        == "SAL-001"
    )

    assert coverage["amount"] == 30000.0

    assert coverage["requires_review"] is False        

