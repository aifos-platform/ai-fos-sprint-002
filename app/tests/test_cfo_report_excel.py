from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook

from app.services.cfo_report_excel import (
    generate_cfo_report_excel,
)


EXPECTED_SHEETS = [
    "Executive Summary",
    "Financial Health",
    "Liquidity",
    "Budget Performance",
    "Funding & Grants",
    "Core Cost Coverage",
    "Current Risks",
    "Forward Risks",
    "Opportunities",
    "CFO Recommendations",
    "Financial Trends",
    "Financial Forecast",
    "Methodology",
]


def _sample_report() -> dict:
    return {
        "version": "1.0",
        "executive_summary": {
            "financial_health": {
                "score": 54,
                "rating": "Watch",
            },
            "liquidity": {
                "cash_runway_months": 11.44,
            },
            "high_current_risk_count": 2,
            "high_forward_risk_count": 1,
            "high_opportunity_count": 1,
            "priority_action_count": 2,
        },
        "financial_health": {
            "score": 54,
            "maximum": 100,
            "rating": "Watch",
            "categories": {
                "liquidity": {
                    "score": 14,
                    "maximum": 20,
                    "reason": "Liquidity remains available.",
                },
                "operating_result": {
                    "score": 8,
                    "maximum": 20,
                    "reason": "Operating deficit requires attention.",
                },
            },
            "metrics": {
                "remaining_requirement": 1000000.0,
                "applied_secured_funding": 600000.0,
                "funding_gap": 400000.0,
                "applied_coverage_percentage": 60.0,
                "matched_requirement_count": 4,
                "unmatched_requirement_count": 2,
            },
        },
        "liquidity": {
            "total_cash": 3283081.45,
            "available_cash": 1971271.45,
            "blocked_cash": 1311810.00,
            "available_account_count": 3,
            "blocked_account_count": 1,
            "unclassified_cash_account_count": 0,
            "total_expenses": 2067000.0,
            "average_monthly_expenses": 172250.0,
            "cash_runway_months": 11.44,
            "reporting_period_start": "2025-09-01",
            "reporting_period_end": "2026-08-31",
            "reporting_month_count": 12,
            "runway_basis": (
                "Trailing operating-expense window."
            ),
        },
        "budget": {
            "executive_summary": {
                "total_budget": 2500000.0,
                "total_actual": 1700000.0,
                "budgeted_actual": 1650000.0,
                "unbudgeted_actual": 50000.0,
                "budget_variance": 850000.0,
                "utilization_percentage": 66.0,
                "over_budget_count": 2,
                "within_budget_count": 10,
                "no_budget_count": 1,
            },
            "organization": {
                "fund_count": 5,
                "donor_count": 8,
                "program_count": 4,
                "project_count": 6,
            },
            "alerts": [
                {
                    "title": "Budget pressure",
                    "severity": "High",
                    "category": "Budget",
                    "reason": "Two lines are over budget.",
                }
            ],
        },
        "funding": {
            "funding_gap": {
                "summary": {
                    "total_needed_budget": 2000000.0,
                    "actual_spending_against_need": 500000.0,
                    "remaining_requirement": 1500000.0,
                    "gross_remaining_secured_funding": 1000000.0,
                    "explicitly_mapped_secured_funding": 900000.0,
                    "applied_secured_funding": 900000.0,
                    "excess_eligible_funding": 0.0,
                    "secured_funding_without_budget_line_allocation": 100000.0,
                    "funding_gap": 600000.0,
                    "applied_coverage_percentage": 60.0,
                    "matched_requirement_count": 5,
                    "unmatched_requirement_count": 2,
                    "fully_funded_requirement_count": 3,
                    "partially_funded_requirement_count": 2,
                    "unfunded_requirement_count": 2,
                    "no_remaining_requirement_count": 1,
                    "requirements_with_dimension_incompatible_funding": 0,
                    "requirements_with_period_ineligible_funding": 1,
                    "requirements_with_period_unknown_funding": 1,
                },
                "methodology": {
                    "eligibility_rule": (
                        "fiscal_year_overlap_required"
                    ),
                    "financial_recalculation_performed": False,
                },
            },
            "grant_diagnostics": {
                "matched_grants": [
                    {"fund": "FR001"},
                    {"fund": "FR002"},
                ],
                "budget_only_grants": [
                    {"fund": "FR003"},
                ],
                "actual_only_grants": [],
            },
        },         

        "core_cost_coverage": {
            "status": "available",
            "summary": {
                "needed_core_cost": 100000.0,
                "direct_grant_coverage": 40000.0,
                "allocated_indirect_recovery": 20000.0,
                "unrestricted_core_funding": 15000.0,
                "remaining_core_cost_gap": 25000.0,
                "core_cost_coverage_percentage": 75.0,
                "available_indirect_recovery": 30000.0,
                "used_indirect_recovery": 12000.0,
            },
            "lines": [
                {
                    "program_code": "CORE",
                    "category_code": "Personnel",
                    "budget_line_code": "SAL-001",
                    "budget_line_name": "Core Salaries",
                    "employee_responsible": "Finance Team",
                    "fiscal_year": 2026,
                    "needed_core_cost": 100000.0,
                    "direct_grant_coverage": 40000.0,
                    "allocated_indirect_recovery": 20000.0,
                    "unrestricted_core_funding": 15000.0,
                    "remaining_core_cost_gap": 25000.0,
                }
            ],
            "controls": {},
        },


        "risks": {
            "current": [
                {
                    "title": "Operating deficit",
                    "severity": "High",
                    "category": "Operating Performance",
                    "evidence": "Expenses exceed revenue.",
                    "recommendation": (
                        "Review operating cost structure."
                    ),
                }
            ],
            "forward": [
                {
                    "title": "Future funding pressure",
                    "severity": "High",
                    "category": "Funding",
                    "evidence": (
                        "Future requirements exceed secured funding."
                    ),
                    "recommendation": (
                        "Prioritize funding pipeline."
                    ),
                }
            ],
        },
        "opportunities": {
            "all": [
                {
                    "title": "Flexible funding opportunity",
                    "priority": "High",
                    "category": "Funding",
                    "evidence": (
                        "Flexible support can strengthen coverage."
                    ),
                    "recommended_action": (
                        "Protect flexible funding allocation."
                    ),
                }
            ],
        },
        "recommendations": {
            "all": [
                {
                    "title": "Protect liquidity",
                    "priority": "Critical",
                    "category": "Liquidity",
                    "evidence": "Runway requires monitoring.",
                    "action": (
                        "Maintain monthly liquidity review."
                    ),
                    "expected_impact": (
                        "Earlier management response."
                    ),
                },
                {
                    "title": "Address budget pressure",
                    "priority": "High",
                    "category": "Budget",
                    "evidence": "Two budget lines are over budget.",
                    "action": "Review budget reallocations.",
                    "expected_impact": (
                        "Improved budget control."
                    ),
                },
            ],
        },
        "trends": {
            "status": "available",
            "basis": "validated_general_ledger",
            "coverage": {
                "month_count": 12,
                "year_count": 2,
            },
            "quality": {
                "status": "sufficient",
            },
            "latest_month_comparison": {
                "current_period": "2026-08",
                "previous_period": "2026-07",
                "revenue": {
                    "current_value": 200000.0,
                    "previous_value": 180000.0,
                    "change_percentage": 11.11,
                    "direction": "increased",
                },
                "expenses": {
                    "current_value": 220000.0,
                    "previous_value": 210000.0,
                    "change_percentage": 4.76,
                    "direction": "increased",
                },
                "net_result": {
                    "current_value": -20000.0,
                    "previous_value": -30000.0,
                    "change_percentage": 33.33,
                    "direction": "improved",
                },
            },
        },
        "forecast": {
            "status": "available",
            "forecast_type": "deterministic",
            "methodology": "historical_run_rate",
            "methodology_description": (
                "Forecast derived from validated historical activity."
            ),
            "forecast_horizon_months": 12,
            "baseline": {
                "monthly_revenue": 200000.0,
                "monthly_expenses": 220000.0,
            },
            "forecast_totals": {
                "revenue": 2400000.0,
                "expenses": 2640000.0,
                "net_result": -240000.0,
            },
            "confidence": {
                "level": "medium",
                "reason": "Based on available historical coverage.",
            },
        },
        "methodology": {
            "source": "validated_financial_model",
            "financial_recalculation_performed": False,
            "hypothetical_scenarios_included": False,
        },
    }


def _load_workbook(
    report: dict | None,
):
    excel_bytes = generate_cfo_report_excel(
        report
    )

    return load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )


def _sheet_values(sheet) -> list:
    values = []

    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                values.append(cell.value)

    return values


def test_generate_cfo_report_excel_returns_valid_xlsx_bytes():
    report = _sample_report()

    excel_bytes = generate_cfo_report_excel(
        report
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # XLSX files are ZIP-based Office Open XML files.
    assert excel_bytes[:2] == b"PK"

    workbook = load_workbook(
        BytesIO(excel_bytes)
    )

    assert workbook is not None


def test_generate_cfo_report_excel_contains_expected_sheets():
    workbook = _load_workbook(
        _sample_report()
    )

    assert workbook.sheetnames == EXPECTED_SHEETS


def test_generate_cfo_report_excel_preserves_financial_values():
    report = _sample_report()

    workbook = _load_workbook(report)

    liquidity_values = _sheet_values(
        workbook["Liquidity"]
    )

    funding_values = _sheet_values(
        workbook["Funding & Grants"]
    )

    assert 1971271.45 in liquidity_values
    assert 1311810.00 in liquidity_values
    assert 11.44 in liquidity_values

    assert 1500000.0 in funding_values
    assert 900000.0 in funding_values
    assert 600000.0 in funding_values
    assert 60.0 in funding_values

def test_generate_cfo_report_excel_includes_core_cost_coverage():
    workbook = _load_workbook(
        _sample_report()
    )

    values = _sheet_values(
        workbook["Core Cost Coverage"]
    )

    assert "Core Cost Coverage" in values
    assert "Needed Core Cost" in values
    assert 100000.0 in values
    assert "Remaining Core Cost Gap" in values
    assert 25000.0 in values

    assert "Budget Line Code" in values
    assert "SAL-001" in values
    assert "Core Salaries" in values
    assert "Finance Team" in values

def test_generate_cfo_report_excel_keeps_current_and_forward_risks_separate():
    workbook = _load_workbook(
        _sample_report()
    )

    current_values = _sheet_values(
        workbook["Current Risks"]
    )

    forward_values = _sheet_values(
        workbook["Forward Risks"]
    )

    assert "Operating deficit" in current_values

    assert (
        "Future funding pressure"
        not in current_values
    )

    assert (
        "Future funding pressure"
        in forward_values
    )

    assert (
        "Operating deficit"
        not in forward_values
    )


def test_generate_cfo_report_excel_preserves_opportunities_and_recommendations():
    workbook = _load_workbook(
        _sample_report()
    )

    opportunity_values = _sheet_values(
        workbook["Opportunities"]
    )

    recommendation_values = _sheet_values(
        workbook["CFO Recommendations"]
    )

    assert (
        "Flexible funding opportunity"
        in opportunity_values
    )

    assert "Protect liquidity" in recommendation_values

    assert (
        "Address budget pressure"
        in recommendation_values
    )


def test_generate_cfo_report_excel_handles_empty_report():
    workbook = _load_workbook({})

    assert workbook.sheetnames == EXPECTED_SHEETS

    assert (
        workbook["Executive Summary"]["A1"].value
        == "CFO Financial Intelligence Report"
    )

    assert (
        workbook["Methodology"]["A1"].value
        == "Methodology & Evidence"
    )


def test_generate_cfo_report_excel_handles_none_report():
    workbook = _load_workbook(None)

    assert workbook.sheetnames == EXPECTED_SHEETS


def test_generate_cfo_report_excel_does_not_mutate_source_report():
    report = _sample_report()

    original = deepcopy(report)

    generate_cfo_report_excel(report)

    assert report == original


def test_generate_cfo_report_excel_contains_evidence_principle():
    workbook = _load_workbook(
        _sample_report()
    )

    methodology_values = [
        str(value)
        for value in _sheet_values(
            workbook["Methodology"]
        )
    ]

    combined = " ".join(
        methodology_values
    )

    assert (
        "does not independently recalculate "
        "financial results"
        in combined
    )


def test_generate_cfo_report_excel_keeps_recommendations_priority_order():
    workbook = _load_workbook(
        _sample_report()
    )

    sheet = workbook[
        "CFO Recommendations"
    ]

    values = [
        cell.value
        for cell in sheet["B"]
        if cell.value is not None
    ]

    critical_position = values.index(
        "Protect liquidity"
    )

    high_position = values.index(
        "Address budget pressure"
    )

    assert critical_position < high_position

def test_cfo_excel_includes_standard_income_statement_sheet():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.cfo_report_excel import (
        generate_cfo_report_excel,
    )

    standard_income_statement = {
        "report_title": "Standard Income Statement",
        "organisation_name": "AI-FOS Test NGO",
        "base_currency": "USD",
        "reporting_period": {
            "start_date": "2026-01-01",
            "end_date": "2026-12-31",
        },
        "sections": [
            {
                "section": "Revenue",
                "lines": [
                    {
                        "account_number": "700001",
                        "account_name": "Test Grant Revenue",
                        "amount": 125000.0,
                    }
                ],
                "total": 125000.0,
            },
            {
                "section": "Expenses",
                "lines": [
                    {
                        "account_number": "600001",
                        "account_name": "Test Personnel Expense",
                        "amount": 80000.0,
                    }
                ],
                "total": 80000.0,
            },
        ],
        "current_period_result": 45000.0,
        "controls": {
            "presentation_only": True,
        },
    }

    excel_bytes = generate_cfo_report_excel(
        {},
        standard_income_statement=standard_income_statement,
        organisation_name="AI-FOS Test NGO",
        base_currency="USD",
    )

    workbook = load_workbook(
        BytesIO(excel_bytes),
        data_only=True,
    )

    assert "Income Statement" in workbook.sheetnames

    worksheet = workbook["Income Statement"]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "Test Grant Revenue" in values
    assert "Test Personnel Expense" in values    