from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


client = TestClient(app)


def test_budget_upload_persists_core_cost_coverage_outputs(
    tmp_path,
    monkeypatch,
):
    financial_model_folder = (
        Path(tmp_path)
        / "financial_model"
    )

    financial_model_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    upload_folder = (
        Path(tmp_path)
        / "uploads"
    )

    upload_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    monkeypatch.setattr(
        "app.main.UPLOAD_FOLDER",
        upload_folder,
    )

    monkeypatch.setattr(
        "app.main.workspace_service.create_workspace",
        lambda **kwargs: {
            "workspace_id": "test-workspace",
            "status": "ready",
            "paths": {
                "financial_model": str(
                    financial_model_folder
                ),
            },
        },
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "save_upload"
        ),
        lambda **kwargs: {
            "status": "saved",
        },
    )

    workbook = Workbook()

    budget_sheet = workbook.active
    budget_sheet.title = "Available Budget"

    budget_sheet.append(
        [
            "Budget Line Code",
            "Budget Line Code Name",
            "Budget Notes",
            "Total Original Budget (USD)",
        ]
    )

    budget_sheet.append(
        [
            "SAL-001",
            "Salary",
            "Secured budget",
            100000,
        ]
    )

    needed_sheet = workbook.create_sheet(
        "Needed Budget"
    )

    needed_sheet.append(
        [
            "Program Code",
            "Category Code",
            "Budget Line Code",
            "Budget Line Name",
            "Employee Responsible",
            "Needed Budget 2026",
        ]
    )

    needed_sheet.append(
        [
            "PROGRAM-001",
            "PERSONNEL",
            "SAL-001",
            "Salary",
            "Employee A",
            100000,
        ]
    )

    coverage_sheet = workbook.create_sheet(
        "Core Cost Coverage"
    )

    coverage_sheet.append(
        [
            "Coverage Type",
            "Fund Code",
            "Budget Line Code",
            "Amount",
        ]
    )

    coverage_sheet.append(
        [
            "direct_grant_coverage",
            "GRANT-001",
            "SAL-001",
            40000,
        ]
    )

    coverage_sheet.append(
        [
            "indirect_recovery_allocation",
            "GRANT-001",
            "SAL-001",
            15000,
        ]
    )

    coverage_sheet.append(
        [
            "unrestricted_core_funding",
            "CORE-001",
            "SAL-001",
            10000,
        ]
    )

    coverage_sheet.append(
        [
            "available_indirect_recovery",
            "GRANT-001",
            None,
            20000,
        ]
    )

    coverage_sheet.append(
        [
            "used_indirect_recovery",
            None,
            "SAL-001",
            12000,
        ]
    )

    workbook_bytes = BytesIO()

    workbook.save(
        workbook_bytes
    )

    workbook_bytes.seek(0)

    response = client.post(
        "/upload",
        files={
            "file": (
                "budget_with_core_cost_coverage.xlsx",
                workbook_bytes.getvalue(),
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
            ),
        },
        data={
            "organisation_id": "acss",
            "organisation_name": "ACSS",
            "base_currency": "USD",
        },
    )

    print(response.status_code, response.text)
    assert response.status_code == 200

    data = response.json()

    assert data["document_type"] == "budget"

    assert (
        data["needed_budget_line_count"]
        == 1
    )

    assert (
        data["core_cost_coverage_line_count"]
        == 5
    )

    assert (
        data["core_cost_coverage_file"]
        is not None
    )

    assert (
        data["core_cost_coverage_intelligence_file"]
        is not None
    )

    assert (
        financial_model_folder
        / "core_cost_coverage.json"
    ).exists()

    assert (
        financial_model_folder
        / "core_cost_coverage_intelligence.json"
    ).exists()
