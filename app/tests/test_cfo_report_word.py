from io import BytesIO

from docx import Document

from app.services.cfo_report_word import generate_cfo_report_word


def test_generate_cfo_report_word_returns_docx_bytes():
    report = {
        "report_type": "cfo_report",
        "version": "1.0",
        "executive_summary": {
            "financial_health": {
                "score": 75,
                "rating": "Moderate",
            },
        },
        "financial_health": {
            "score": 75,
            "rating": "Moderate",
        },
        "liquidity": {
            "cash_runway_months": 10.5,
        },
        "budget": {
            "status": "available",
        },
        "funding": {
            "funding_gap": {
                "summary": {
                    "funding_gap": 500000.0,
                },
            },
        },
        "core_cost_coverage": {
            "status": "available",
            "summary": {
                "needed_core_cost": 100000.0,
                "direct_grant_coverage": 40000.0,
                "allocated_indirect_recovery": 20000.0,
                "unrestricted_core_funding": 15000.0,
                "remaining_core_cost_gap": 25000.0,
                "core_cost_coverage_percentage": 75.0,
                "available_indirect_recovery": 30000.0,
                "used_indirect_recovery": 12000.0,
            },
        },
        "risks": {},
        "opportunities": {},
        "recommendations": {},
        "trends": {},
        "forecast": {},
        "methodology": {},
    }

    word_bytes = generate_cfo_report_word(report)

    assert isinstance(word_bytes, bytes)
    assert len(word_bytes) > 0
    assert word_bytes[:2] == b"PK"

    document = Document(BytesIO(word_bytes))

    paragraph_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "CFO Financial Intelligence Report" in paragraph_text
    assert "Core Cost Coverage" in paragraph_text

def test_generate_cfo_report_word_handles_empty_report():
    word_bytes = generate_cfo_report_word({})

    assert isinstance(word_bytes, bytes)
    assert len(word_bytes) > 0
    assert word_bytes[:2] == b"PK"

    document = Document(BytesIO(word_bytes))

    paragraph_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "CFO Financial Intelligence Report" in paragraph_text 

def test_generate_cfo_report_word_includes_professional_sections():
    report = {
        "executive_summary": {},
        "financial_health": {},
        "liquidity": {},
        "budget": {},
        "funding": {},
        "core_cost_coverage": {},
        "risks": {},
        "opportunities": {},
        "recommendations": {},
        "trends": {},
        "forecast": {},
        "methodology": {},
    }

    word_bytes = generate_cfo_report_word(report)

    document = Document(BytesIO(word_bytes))

    paragraph_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    expected_sections = [
        "Executive Summary",
        "Financial Health",
        "Liquidity",
        "Budget Performance",
        "Funding & Grants",
        "Core Cost Coverage",
        "Current & Forward Risks",
        "Opportunities",
        "CFO Recommendations",
        "Trends",
        "Forecast",
        "Methodology",
    ]

    for section in expected_sections:
        assert section in paragraph_text

def test_generate_cfo_report_word_includes_executive_summary_values():
    report = {
        "executive_summary": {
            "financial_health": {
                "score": 75,
                "maximum": 100,
                "rating": "Moderate",
            },
            "liquidity": {
                "cash_runway_months": 10.5,
            },
            "high_current_risk_count": 2,
            "high_forward_risk_count": 3,
            "high_opportunity_count": 4,
            "priority_action_count": 5,
        },
        "financial_health": {
            "score": 75,
            "maximum": 100,
            "rating": "Moderate",
        },
        "liquidity": {
            "cash_runway_months": 10.5,
        },
        "funding": {
            "funding_gap": {
                "summary": {
                    "funding_gap": 500000.0,
                    "applied_coverage_percentage": 60.0,
                },
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)

    document = Document(BytesIO(word_bytes))

    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    assert "Financial Health" in table_text
    assert "75 / 100" in table_text
    assert "Moderate" in table_text
    assert "Cash Runway" in table_text
    assert "10.5" in table_text
    assert "Funding Gap" in table_text
    assert "500,000.00" in table_text
    assert "60.0%" in table_text  

def test_generate_cfo_report_word_includes_financial_health_details():
    report = {
        "financial_health": {
            "score": 75,
            "maximum": 100,
            "rating": "Moderate",
            "categories": {
                "liquidity": {
                    "score": 18,
                    "maximum": 25,
                    "reason": "Liquidity remains adequate.",
                },
                "sustainability": {
                    "score": 15,
                    "maximum": 25,
                    "reason": "Funding coverage requires attention.",
                },
            },
            "metrics": {
                "remaining_requirement": 800000.0,
                "applied_secured_funding": 500000.0,
                "funding_gap": 300000.0,
                "applied_coverage_percentage": 62.5,
                "matched_requirement_count": 5,
                "unmatched_requirement_count": 2,
                "actual_only_grant_count": 1,
                "budget_only_grant_count": 3,
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Financial Health" in document_text
    assert "75 / 100" in document_text
    assert "Moderate" in document_text

    assert "Liquidity" in document_text
    assert "18" in document_text
    assert "25" in document_text
    assert "Liquidity remains adequate." in document_text

    assert "Sustainability" in document_text
    assert "Funding coverage requires attention." in document_text

    assert "Key Financial Health Metrics" in document_text
    assert "remaining_requirement" in document_text
    assert "800000.0" in document_text
    assert "applied_coverage_percentage" in document_text
    assert "62.5" in document_text
    assert "budget_only_grant_count" in document_text
    assert "3" in document_text   

def test_generate_cfo_report_word_includes_liquidity_details():
    report = {
        "liquidity": {
            "available_cash": 1971271.45,
            "blocked_cash": 1311810.00,
            "cash_runway_months": 11.44,
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Available Cash" in document_text
    assert "1,971,271.45" in document_text
    assert "Validated available liquidity" in document_text

    assert "Blocked Cash" in document_text
    assert "1,311,810.00" in document_text
    assert "Restricted / blocked liquidity" in document_text

    assert "Cash Runway" in document_text
    assert "11.44 months" in document_text
    assert "Operating-expense coverage" in document_text    

def test_generate_cfo_report_word_includes_budget_performance():
    report = {
        "budget": {
            "executive_summary": {
                "total_budget": 1000000.0,
                "total_actual": 750000.0,
                "budgeted_actual": 700000.0,
                "unbudgeted_actual": 50000.0,
                "budget_variance": 300000.0,
                "total_variance": 250000.0,
                "overall_variance_including_unbudgeted": 250000.0,
                "utilization_percentage": 75.0,
                "budget_line_count": 20,
                "line_count": 22,
                "over_budget_count": 3,
                "within_budget_count": 15,
                "no_budget_count": 4,
            },
            "organization": {
                "fund_count": 8,
                "donor_count": 6,
                "program_count": 5,
                "project_count": 10,
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Budget Performance" in document_text

    assert "Total Budget" in document_text
    assert "1,000,000.00" in document_text

    assert "Total Actual" in document_text
    assert "750,000.00" in document_text

    assert "utilization_percentage" in document_text
    assert "75.0" in document_text

    assert "over_budget_count" in document_text
    assert "3" in document_text

    assert "within_budget_count" in document_text
    assert "15" in document_text

    assert "Organization Scope" in document_text
    assert "fund_count" in document_text
    assert "8" in document_text
    assert "donor_count" in document_text
    assert "6" in document_text
    assert "program_count" in document_text
    assert "5" in document_text
    assert "project_count" in document_text
    assert "10" in document_text

def test_generate_cfo_report_word_includes_funding_details():
    report = {
        "funding": {
            "funding_gap": {
                "summary": {
                    "remaining_requirement": 800000.0,
                    "applied_secured_funding": 500000.0,
                    "excess_eligible_funding": 25000.0,
                    "secured_funding_without_budget_line_allocation": 40000.0,
                    "funding_gap": 300000.0,
                    "applied_coverage_percentage": 62.5,
                    "matched_requirement_count": 5,
                    "unmatched_requirement_count": 2,
                    "fully_funded_requirement_count": 3,
                    "partially_funded_requirement_count": 2,
                    "unfunded_requirement_count": 2,
                    "no_remaining_requirement_count": 1,
                    "requirements_with_dimension_incompatible_funding": 1,
                    "requirements_with_period_ineligible_funding": 2,
                    "requirements_with_period_unknown_funding": 1,
                    "diagnostic_exposure_note": (
                        "Some secured funding requires review."
                    ),
                },
            },
            "grant_diagnostics": {
                "matched_grants": [
                    {"fund_code": "FR001"},
                    {"fund_code": "FR002"},
                ],
                "budget_only_grants": [
                    {"fund_code": "FR003"},
                ],
                "actual_only_grants": [
                    {"fund_code": "FR004"},
                    {"fund_code": "FR005"},
                    {"fund_code": "FR006"},
                ],
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Funding & Grants" in document_text

    assert "Remaining Requirement" in document_text
    assert "800,000.00" in document_text

    assert "Applied Secured Funding" in document_text
    assert "500,000.00" in document_text

    assert "Funding Gap" in document_text
    assert "300,000.00" in document_text

    assert "62.5% coverage" in document_text

    assert "excess_eligible_funding" in document_text
    assert "secured_funding_without_budget_line_allocation" in document_text

    assert "Funding Diagnostic Note" in document_text
    assert "Some secured funding requires review." in document_text

    assert "Grant Diagnostics" in document_text
    assert "matched_grants" in document_text
    assert "2" in document_text
    assert "budget_only_grants" in document_text
    assert "1" in document_text
    assert "actual_only_grants" in document_text
    assert "3" in document_text 

def test_generate_cfo_report_word_includes_core_cost_coverage():
    report = {
        "core_cost_coverage": {
            "status": "available",
            "summary": {
                "needed_core_cost": 1000000.0,
                "direct_grant_coverage": 400000.0,
                "allocated_indirect_recovery": 150000.0,
                "unrestricted_core_funding": 100000.0,
                "remaining_core_cost_gap": 350000.0,
                "core_cost_coverage_percentage": 65.0,
                "available_indirect_recovery": 250000.0,
                "used_indirect_recovery": 120000.0,
            },
            "lines": [
                {
                    "program_code": "PRG001",
                    "category_code": "Personnel",
                    "budget_line_code": "BL001",
                    "budget_line_name": "Program Manager",
                    "employee_responsible": "Employee A",
                    "fiscal_year": 2026,
                    "needed_core_cost": 100000.0,
                    "direct_grant_coverage": 40000.0,
                    "allocated_indirect_recovery": 15000.0,
                    "unrestricted_core_funding": 10000.0,
                    "remaining_core_cost_gap": 35000.0,
                },
            ],
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Core Cost Coverage" in document_text

    assert "Needed Core Cost" in document_text
    assert "1,000,000.00" in document_text
    assert "Validated core cost requirement" in document_text

    assert "Direct Grant Coverage" in document_text
    assert "400,000.00" in document_text

    assert "Allocated Indirect Recovery" in document_text
    assert "150,000.00" in document_text

    assert "Unrestricted / Core Funding" in document_text
    assert "100,000.00" in document_text

    assert "Remaining Core Cost Gap" in document_text
    assert "350,000.00" in document_text

    assert "65.0%" in document_text

    assert "Available Indirect Recovery" in document_text
    assert "250,000.00" in document_text

    assert "Used / Charged Indirect" in document_text
    assert "120,000.00" in document_text

    assert "Core Cost Coverage Lines" in document_text
    assert "BL001" in document_text
    assert "Program Manager" in document_text
    assert "35,000.00" in document_text 

def test_generate_cfo_report_word_includes_current_and_forward_risks():
    report = {
        "risks": {
            "current": [
                {
                    "title": "Operating Deficit",
                    "severity": "high",
                    "category": "financial_sustainability",
                    "evidence": "Operating expenses exceed operating revenue.",
                    "recommendation": "Strengthen unrestricted funding coverage.",
                    "action": "Review the operating cost base.",
                },
            ],
            "forward": [
                {
                    "title": "Future Funding Gap",
                    "severity": "high",
                    "category": "forward_funding",
                    "evidence": "Future requirements exceed secured funding.",
                    "recommendation": "Prioritize pipeline conversion.",
                    "action": "Review upcoming grant opportunities.",
                },
            ],
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Current & Forward Risks" in document_text

    assert "Current Financial Risks" in document_text
    assert "Operating Deficit" in document_text
    assert "high" in document_text
    assert "financial_sustainability" in document_text
    assert "Operating expenses exceed operating revenue." in document_text
    assert "Strengthen unrestricted funding coverage." in document_text
    assert "Review the operating cost base." in document_text

    assert "Forward-Looking Risks" in document_text
    assert "Future Funding Gap" in document_text
    assert "forward_funding" in document_text
    assert "Future requirements exceed secured funding." in document_text
    assert "Prioritize pipeline conversion." in document_text
    assert "Review upcoming grant opportunities." in document_text

def test_generate_cfo_report_word_includes_opportunities():
    report = {
        "opportunities": {
            "all": [
                {
                    "title": "Strengthen Core Funding",
                    "priority": "high",
                    "category": "funding",
                    "evidence": "Unrestricted funding can improve financial resilience.",
                    "recommended_action": "Expand the unrestricted funding pipeline.",
                    "action": "Engage priority prospective donors.",
                },
            ],
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Opportunities" in document_text
    assert "Strengthen Core Funding" in document_text
    assert "high" in document_text
    assert "funding" in document_text
    assert "Unrestricted funding can improve financial resilience." in document_text
    assert "Expand the unrestricted funding pipeline." in document_text
    assert "Engage priority prospective donors." in document_text   

def test_generate_cfo_report_word_includes_cfo_recommendations():
    report = {
        "recommendations": {
            "all": [
                {
                    "title": "Strengthen Funding Pipeline",
                    "priority": "high",
                    "category": "funding",
                    "evidence": "Future funding requirements exceed secured funding.",
                    "reason": "Additional funding is required to reduce future exposure.",
                    "expected_impact": "Improved future funding coverage.",
                    "action": "Prioritize high-probability donor opportunities.",
                },
            ],
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "CFO Recommendations" in document_text
    assert "Strengthen Funding Pipeline" in document_text
    assert "high" in document_text
    assert "funding" in document_text
    assert "Future funding requirements exceed secured funding." in document_text
    assert "Additional funding is required to reduce future exposure." in document_text
    assert "Improved future funding coverage." in document_text
    assert "Prioritize high-probability donor opportunities." in document_text

def test_generate_cfo_report_word_includes_trends():
    report = {
        "trends": {
            "coverage": {
                "months_available": 12,
                "years_available": 2,
            },
            "quality": {
                "status": "good",
                "missing_periods": 0,
            },
            "latest_month_comparison": {
                "current_period": "2026-08",
                "previous_period": "2026-07",
                "revenue": {
                    "current_value": 125000.0,
                    "previous_value": 100000.0,
                    "change_percentage": 25.0,
                    "direction": "up",
                },
                "expenses": {
                    "current_value": 90000.0,
                    "previous_value": 85000.0,
                    "change_percentage": 5.88,
                    "direction": "up",
                },
                "net_result": {
                    "current_value": 35000.0,
                    "previous_value": 15000.0,
                    "change_percentage": 133.33,
                    "direction": "up",
                },
            },
            "latest_year_comparison": {
                "current_period": "2026",
                "previous_period": "2025",
                "revenue": {
                    "current_value": 1500000.0,
                    "previous_value": 1400000.0,
                    "change_percentage": 7.14,
                    "direction": "up",
                },
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Trends" in document_text
    assert "Trend Coverage" in document_text
    assert "months_available" in document_text
    assert "12" in document_text

    assert "Data Quality" in document_text
    assert "status" in document_text
    assert "good" in document_text

    assert "Latest Month Comparison" in document_text
    assert "2026-08" in document_text
    assert "2026-07" in document_text
    assert "Revenue" in document_text
    assert "125,000.00" in document_text
    assert "100,000.00" in document_text
    assert "25.0%" in document_text
    assert "up" in document_text

    assert "Latest Year Comparison" in document_text
    assert "2026" in document_text
    assert "2025" in document_text
    assert "1,500,000.00" in document_text
    assert "1,400,000.00" in document_text
    assert "7.1%" in document_text

def test_generate_cfo_report_word_includes_forecast():
    report = {
        "forecast": {
            "baseline": {
                "forecast_horizon_months": 12,
                "method": "deterministic",
            },
            "forecast_totals": {
                "revenue": 1500000.0,
                "expenses": 1250000.0,
                "net_result": 250000.0,
            },
            "confidence": {
                "level": "medium",
                "reason": "Based on validated historical financial data.",
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Forecast" in document_text

    assert "Baseline" in document_text
    assert "forecast_horizon_months" in document_text
    assert "12" in document_text
    assert "deterministic" in document_text

    assert "Forecast Revenue" in document_text
    assert "1,500,000.00" in document_text

    assert "Forecast Expenses" in document_text
    assert "1,250,000.00" in document_text

    assert "Forecast Net Result" in document_text
    assert "250,000.00" in document_text

    assert "Forecast Confidence" in document_text
    assert "medium" in document_text
    assert "Based on validated historical financial data." in document_text

def test_generate_cfo_report_word_includes_methodology():
    report = {
        "methodology": {
            "calculation_policy": (
                "The CFO Report consumes validated AI-FOS financial "
                "intelligence outputs and does not recalculate financial results."
            ),
            "evidence_policy": (
                "Report sections remain traceable to validated AI-FOS "
                "financial model artifacts."
            ),
        },
        "funding": {
            "funding_gap": {
                "methodology": {
                    "eligibility_rule": "fiscal_year_overlap_required",
                    "allocation_policy": "explicit_mapping_only",
                },
            },
        },
    }

    word_bytes = generate_cfo_report_word(report)
    document = Document(BytesIO(word_bytes))

    document_text = "\n".join(
        [
            *(
                paragraph.text
                for paragraph in document.paragraphs
            ),
            *(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ),
        ]
    )

    assert "Methodology" in document_text
    assert "calculation_policy" in document_text
    assert "does not recalculate financial results" in document_text
    assert "evidence_policy" in document_text
    assert "financial model artifacts" in document_text

    assert "Funding Gap Methodology" in document_text
    assert "eligibility_rule" in document_text
    assert "fiscal_year_overlap_required" in document_text
    assert "allocation_policy" in document_text
    assert "explicit_mapping_only" in document_text

    assert "Evidence Principle" in document_text
    assert "consumes validated financial model outputs" in document_text                                             