from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.organization import Organization
from app.services.financial_model_service import FinancialModelService


client = TestClient(app)


def _workspace(financial_model_folder):
    return {
        "organisation_id": "test-org",
        "paths": {
            "financial_model": str(financial_model_folder),
        },
    }


def _organisation_metadata():
    return {
        "organisation_id": "test-org",
        "name": "Test NGO",
        "base_currency": "USD",
    }


def _standard_balance_sheet():
    return {
        "status": "available",
        "report_type": "balance_sheet",
        "title": "Balance Sheet",
        "reporting_basis": "as_of_date",
        "sections": [
            {
                "section": "Assets",
                "lines": [
                    {
                        "account_number": "201100",
                        "account_name": "Cash",
                        "parent_account": "201000",
                        "level": 2,
                        "financial_category": "Asset",
                        "financial_subcategory": "Cash",
                        "normal_balance": "Debit",
                        "amount": 1500.0,
                    },
                ],
                "total": 1500.0,
            },
            {
                "section": "Liabilities",
                "lines": [
                    {
                        "account_number": "401100",
                        "account_name": "Accounts Payable",
                        "parent_account": "401000",
                        "level": 2,
                        "financial_category": "Liability",
                        "financial_subcategory": "Payables",
                        "normal_balance": "Credit",
                        "amount": 900.0,
                    },
                ],
                "total": 900.0,
            },
            {
                "section": "Equity",
                "lines": [
                    {
                        "account_number": "101100",
                        "account_name": "Equity",
                        "parent_account": "101000",
                        "level": 2,
                        "financial_category": "Equity",
                        "financial_subcategory": "Capital and Equity",
                        "normal_balance": "Credit",
                        "amount": 400.0,
                    },
                ],
                "total": 400.0,
            },
        ],
        "totals": {
            "assets": 1500.0,
            "liabilities": 900.0,
            "reported_equity": 400.0,
            "current_period_result": 200.0,
            "equity": 600.0,
            "liabilities_and_equity": 1500.0,
            "difference": 0.0,
        },
        "reconciliation": {
            "status": "reconciled",
            "difference": 0.0,
        },
        "controls": {
            "reporting_layer_only": True,
            "validated_balance_sheet_preserved": True,
            "financial_recalculation_performed": False,
            "missing_accounts_not_invented": True,
            "reconciliation_required": True,
        },
    }


def _minimal_cfo_report():
    return {
        "status": "available",
        "executive_summary": {},
        "financial_health": {},
        "liquidity": {},
        "budget": {},
        "funding": {},
        "risks": {},
        "opportunities": {},
        "recommendations": {},
        "trends": {},
        "forecast": {},
        "methodology": {},
    }


def test_cfo_excel_download_includes_persisted_standard_balance_sheet(
    tmp_path,
    monkeypatch,
):
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organisation_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "cfo_report.json",
        _minimal_cfo_report(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_balance_sheet.json",
        _standard_balance_sheet(),
    )

    response = client.get(
        "/reports/test-org/balance-sheet/excel"
    )

    assert response.status_code == 200

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    assert "Balance Sheet" in workbook.sheetnames

    worksheet = workbook["Balance Sheet"]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "ASSETS" in values
    assert "LIABILITIES" in values
    assert "EQUITY" in values
    assert "Cash" in values
    assert "Accounts Payable" in values

    assert 1500.0 in values
    assert 900.0 in values
    assert 400.0 in values


def test_cfo_excel_balance_sheet_download_does_not_recalculate_financials(
    tmp_path,
    monkeypatch,
):
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organisation_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "cfo_report.json",
        _minimal_cfo_report(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_balance_sheet.json",
        _standard_balance_sheet(),
    )

    def forbidden_recalculation(
        *args,
        **kwargs,
    ):
        raise AssertionError(
            "CFO Excel Balance Sheet download must not "
            "recalculate financials."
        )

    monkeypatch.setattr(
        Organization,
        "process_financials",
        forbidden_recalculation,
    )

    response = client.get(
        "/reports/test-org/balance-sheet/excel"
    )


    assert response.status_code == 200

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    assert "Balance Sheet" in workbook.sheetnames


def test_cfo_excel_balance_sheet_preserves_persisted_report(
    tmp_path,
    monkeypatch,
):
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organisation_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "cfo_report.json",
        _minimal_cfo_report(),
    )

    report = _standard_balance_sheet()

    FinancialModelService.save_json(
        tmp_path,
        "standard_balance_sheet.json",
        report,
    )

    before = FinancialModelService.load_json(
        tmp_path,
        "standard_balance_sheet.json",
    )

    response = client.get(
        "/reports/test-org/balance-sheet/excel"
    )

    assert response.status_code == 200

    after = FinancialModelService.load_json(
        tmp_path,
        "standard_balance_sheet.json",
    )

    assert after == before
    assert after["totals"]["assets"] == 1500.0
    assert after["totals"]["equity"] == 600.0
    assert after["reconciliation"]["status"] == "reconciled"
    assert (
        after["controls"]["financial_recalculation_performed"]
        is False
    )