from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO

from app.main import app
from app.organization import Organization
from app.services.financial_model_service import (
    FinancialModelService,
)


client = TestClient(app)


def _workspace(
    financial_model_folder: Path,
) -> dict:
    return {
        "paths": {
            "financial_model": str(
                financial_model_folder
            ),
        },
    }


def _standard_income_statement() -> dict:
    return {
        "status": "available",
        "report_type": "standard_income_statement",
        "totals": {
            "revenue": 1000.0,
            "expenses": 500.0,
            "net_result": 500.0,
        },
        "sections": [],
        "reconciliation": {
            "reconciled": True,
        },
        "controls": {
            "reporting_layer_only": True,
            "financial_recalculation_performed": False,
        },
    }


def _organization_metadata() -> dict:
    return {
        "id": "test-org",
        "name": "Test NGO",
        "base_currency": "EUR",
        "active": True,
    }


def test_standard_income_statement_excel_api_returns_workbook(
    monkeypatch,
    tmp_path: Path,
):
    report = _standard_income_statement()
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organization_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_income_statement.json",
        report,
    )

    response = client.get(
        "/reports/test-org/income-statement/excel"
    )

    assert response.status_code == 200

    assert (
        response.headers["content-type"]
        == (
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    assert (
        'filename="AI-FOS_test-org_Income_Statement.xlsx"'
        in response.headers["content-disposition"]
    )

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    assert workbook.sheetnames == [
        "Income Statement"
    ]

    worksheet = workbook[
        "Income Statement"
    ]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "Test NGO" in values
    assert "Income Statement" in values


def test_standard_income_statement_excel_api_uses_registry_metadata(
    monkeypatch,
    tmp_path: Path,
):
    report = _standard_income_statement()
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: {
            "id": "test-org",
            "name": "Global Relief NGO",
            "base_currency": "GBP",
            "active": True,
        },
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_income_statement.json",
        report,
    )

    response = client.get(
        "/reports/test-org/income-statement/excel"
    )

    assert response.status_code == 200

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    worksheet = workbook[
        "Income Statement"
    ]

    values = [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]

    assert "Global Relief NGO" in values

    currency_values = [
        str(value)
        for value in values
        if value is not None
    ]

    assert any(
        "GBP" in value
        for value in currency_values
    )


def test_standard_income_statement_excel_api_missing_workspace_returns_404(
    monkeypatch,
):
    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: None,
    )

    response = client.get(
        "/reports/missing-org/income-statement/excel"
    )

    assert response.status_code == 404

    assert (
        response.json()["detail"]
        == "No workspace found for this organization."
    )


def test_standard_income_statement_excel_api_missing_report_returns_404(
    monkeypatch,
    tmp_path: Path,
):
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    response = client.get(
        "/reports/test-org/income-statement/excel"
    )

    assert response.status_code == 404

    assert (
        "Standard Income Statement is not available yet"
        in response.json()["detail"]
    )


def test_standard_income_statement_excel_api_does_not_recalculate_financials(
    monkeypatch,
    tmp_path: Path,
):
    report = _standard_income_statement()
    workspace = _workspace(tmp_path)

    monkeypatch.setattr(
        "app.main.workspace_service.get_workspace_by_organisation",
        lambda organisation_id: workspace,
    )

    monkeypatch.setattr(
        "app.main.organization_registry_service.get_organization",
        lambda organisation_id: _organization_metadata(),
    )

    FinancialModelService.save_json(
        tmp_path,
        "standard_income_statement.json",
        report,
    )

    def forbidden_recalculation(
        *args,
        **kwargs,
    ):
        raise AssertionError(
            "Income Statement Excel API must not "
            "recalculate financials."
        )

    monkeypatch.setattr(
        Organization,
        "process_financials",
        forbidden_recalculation,
    )

    response = client.get(
        "/reports/test-org/income-statement/excel"
    )

    assert response.status_code == 200

    workbook = load_workbook(
        BytesIO(response.content),
        data_only=False,
    )

    assert workbook.sheetnames == [
        "Income Statement"
    ]