from io import BytesIO

from openpyxl import load_workbook

from app.services.standard_balance_sheet_excel import (
    generate_standard_balance_sheet_excel,
)


def _standard_balance_sheet():
    return {
        "status": "available",
        "report_type": "balance_sheet",
        "title": "Balance Sheet",
        "reporting_basis": "as_of_date",
        "sections": [
            {
                "section": "Assets",
                "lines": [
                    {
                        "account_number": "201100",
                        "account_name": "Cash",
                        "parent_account": "201000",
                        "level": 2,
                        "financial_category": "Asset",
                        "financial_subcategory": "Cash",
                        "normal_balance": "Debit",
                        "amount": 1500.0,
                    },
                ],
                "total": 1500.0,
            },
            {
                "section": "Liabilities",
                "lines": [
                    {
                        "account_number": "401100",
                        "account_name": "Accounts Payable",
                        "parent_account": "401000",
                        "level": 2,
                        "financial_category": "Liability",
                        "financial_subcategory": "Payables",
                        "normal_balance": "Credit",
                        "amount": 900.0,
                    },
                ],
                "total": 900.0,
            },
            {
                "section": "Equity / Net Assets",
                "lines": [
                    {
                        "account_number": "101100",
                        "account_name": "Equity",
                        "parent_account": "101000",
                        "level": 2,
                        "financial_category": "Equity",
                        "financial_subcategory": "Equity",
                        "normal_balance": "Credit",
                        "amount": 400.0,
                    },
                ],
                "reported_equity": 400.0,
                "current_period_result": 200.0,
                "total": 600.0,
            },
        ],
        "totals": {
            "assets": 1500.0,
            "liabilities": 900.0,
            "reported_equity": 400.0,
            "current_period_result": 200.0,
            "equity": 600.0,
            "liabilities_and_equity": 1500.0,
            "difference": 0.0,
        },
        "validated_totals": {
            "assets": 1500.0,
            "liabilities": 900.0,
            "reported_equity": 400.0,
            "current_period_result": 200.0,
            "equity": 600.0,
            "liabilities_and_equity": 1500.0,
            "difference": 0.0,
        },
        "reconciliation": {
            "reconciled": True,
        },
        "controls": {
            "reporting_layer_only": True,
            "validated_balance_sheet_preserved": True,
            "financial_recalculation_performed": False,
            "missing_accounts_not_invented": True,
            "reconciliation_required": True,
        },
    }


def _load_workbook():
    excel_bytes = generate_standard_balance_sheet_excel(
        _standard_balance_sheet(),
        organisation_name="Test NGO",
        base_currency="USD",
    )

    return load_workbook(
        BytesIO(excel_bytes),
        data_only=False,
    )


def _sheet_values(worksheet):
    return [
        [
            cell.value
            for cell in row
        ]
        for row in worksheet.iter_rows()
    ]


def test_standard_balance_sheet_excel_creates_balance_sheet():
    workbook = _load_workbook()

    assert workbook.sheetnames == ["Balance Sheet"]


def test_standard_balance_sheet_excel_contains_organisation():
    workbook = _load_workbook()
    worksheet = workbook["Balance Sheet"]

    assert worksheet["A1"].value == "Test NGO"
    assert worksheet["A2"].value == "Balance Sheet"


def test_standard_balance_sheet_excel_contains_account_lines():
    workbook = _load_workbook()
    worksheet = workbook["Balance Sheet"]

    values = _sheet_values(worksheet)

    flattened = [
        value
        for row in values
        for value in row
    ]

    assert "201100" in flattened
    assert "Cash" in flattened

    assert "401100" in flattened
    assert "Accounts Payable" in flattened

    assert "101100" in flattened
    assert "Equity" in flattened


def test_standard_balance_sheet_excel_contains_section_totals():
    workbook = _load_workbook()
    worksheet = workbook["Balance Sheet"]

    values = _sheet_values(worksheet)

    labels = [
        row[0]
        for row in values
        if row
    ]

    assert "Total Assets" in labels
    assert "Total Liabilities" in labels

    assert (
        "Total Reported Equity / Net Assets"
        in labels
    )

    assert (
        "Current Period Surplus / (Deficit)"
        in labels
    )

    assert (
        "Adjusted Equity / Net Assets"
        in labels
    )


def test_standard_balance_sheet_excel_preserves_validated_values():
    workbook = _load_workbook()
    worksheet = workbook["Balance Sheet"]

    values = _sheet_values(worksheet)

    rows_by_label = {
        row[0]: row
        for row in values
        if row and row[0] is not None
    }

    assert rows_by_label["Total Assets"][3] == 1500.0
    assert rows_by_label["Total Liabilities"][3] == 900.0

    assert (
        rows_by_label[
            "Total Reported Equity / Net Assets"
        ][3]
        == 400.0
    )

    assert (
        rows_by_label[
            "Current Period Surplus / (Deficit)"
        ][3]
        == 200.0
    )

    assert (
        rows_by_label[
            "Adjusted Equity / Net Assets"
        ][3]
        == 600.0
    )

    assert (
        rows_by_label[
            "TOTAL LIABILITIES & EQUITY"
        ][3]
        == 1500.0
    )


def test_standard_balance_sheet_excel_shows_reconciliation():
    workbook = _load_workbook()
    worksheet = workbook["Balance Sheet"]

    values = _sheet_values(worksheet)

    flattened = [
        str(value)
        for row in values
        for value in row
        if value is not None
    ]

    assert (
        "BALANCE SHEET DIFFERENCE — RECONCILED"
        in flattened
    )

    assert any(
        "Reconciled to validated Balance Sheet"
        in value
        for value in flattened
    )


def test_standard_balance_sheet_excel_has_no_formulas():
    workbook = _load_workbook()
    worksheet = workbook["Balance Sheet"]

    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                assert not cell.value.startswith("=")


def test_standard_balance_sheet_excel_does_not_mutate_report():
    report = _standard_balance_sheet()

    original_assets = report["totals"]["assets"]
    original_equity = report["totals"]["equity"]
    original_lines = list(
        report["sections"][0]["lines"]
    )

    generate_standard_balance_sheet_excel(
        report,
        organisation_name="Test NGO",
        base_currency="USD",
    )

    assert report["totals"]["assets"] == original_assets
    assert report["totals"]["equity"] == original_equity

    assert (
        report["sections"][0]["lines"]
        == original_lines
    )