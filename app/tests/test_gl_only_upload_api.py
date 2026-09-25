import json
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.organization import Organization


client = TestClient(app)


def test_gl_only_upload_generates_and_persists_dim_account(
    tmp_path,
    monkeypatch,
):
    upload_folder = (
        Path(tmp_path)
        / "uploads"
    )

    financial_model_folder = (
        Path(tmp_path)
        / "financial_model"
    )

    ai_knowledge_folder = (
        Path(tmp_path)
        / "ai_knowledge"
    )

    upload_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    financial_model_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    ai_knowledge_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    monkeypatch.setattr(
        "app.main.UPLOAD_FOLDER",
        upload_folder,
    )

    monkeypatch.setattr(
        "app.main.organization",
        Organization(),
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "create_workspace"
        ),
        lambda **kwargs: {
            "workspace_id": "test-gl-only-workspace",
            "status": "ready",
            "paths": {
                "financial_model": str(
                    financial_model_folder
                ),
                "ai_knowledge": str(
                    ai_knowledge_folder
                ),
            },
        },
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "save_upload"
        ),
        lambda **kwargs: {
            "status": "saved",
        },
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "General Ledger"

    sheet.append(
        [
            "Posting Date",
            "G/L Account No.",
            "G/L Account Name",
            "Debit Amount (LCY)",
            "Credit Amount (LCY)",
        ]
    )

    sheet.append(
        [
            "2026-01-15",
            "610001",
            "Salaries",
            100.0,
            0.0,
        ]
    )

    sheet.append(
        [
            "2026-01-15",
            "250001",
            "Computer Equipment",
            0.0,
            100.0,
        ]
    )

    workbook_bytes = BytesIO()

    workbook.save(
        workbook_bytes
    )

    workbook_bytes.seek(0)

    response = client.post(
        "/upload",
        files={
            "file": (
                "gl_only.xlsx",
                workbook_bytes.getvalue(),
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
            ),
        },
        data={
            "organisation_id": "gl-only-test",
            "organisation_name": "GL Only Test",
            "base_currency": "USD",
        },
    )

    assert response.status_code == 200

    data = response.json()

    assert data["document_type"] == "general_ledger"

    dim_account_file = (
        financial_model_folder
        / "dim_account.json"
    )

    assert dim_account_file.exists()

    with dim_account_file.open(
        "r",
        encoding="utf-8",
    ) as file:
        persisted_accounts = json.load(file)

    assert len(persisted_accounts) == 2

    persisted_by_number = {
        str(account["account_number"]): account
        for account in persisted_accounts
    }

    assert set(persisted_by_number) == {
        "610001",
        "250001",
    }

    assert (
        persisted_by_number["610001"]["account_name"]
        == "Salaries"
    )

    assert (
        persisted_by_number["250001"]["account_name"]
        == "Computer Equipment"
    )

    for account in persisted_accounts:
        assert account["account_type"] == "Posting"
        assert account["is_posting_account"] is True

def test_gl_upload_preserves_existing_saved_dim_account(
    tmp_path,
    monkeypatch,
):
    upload_folder = (
        Path(tmp_path)
        / "uploads"
    )

    financial_model_folder = (
        Path(tmp_path)
        / "financial_model"
    )

    ai_knowledge_folder = (
        Path(tmp_path)
        / "ai_knowledge"
    )

    upload_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    financial_model_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    ai_knowledge_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    existing_accounts = [
        {
            "account_number": "610001",
            "account_name": "Existing Salaries Account",
            "account_type": "Posting",
            "is_posting_account": True,
            "financial_category": "Expense",
        },
        {
            "account_number": "250001",
            "account_name": "Existing Asset Account",
            "account_type": "Posting",
            "is_posting_account": True,
            "financial_category": "Asset",
        },
    ]

    dim_account_file = (
        financial_model_folder
        / "dim_account.json"
    )

    with dim_account_file.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            existing_accounts,
            file,
            indent=2,
        )

    monkeypatch.setattr(
        "app.main.UPLOAD_FOLDER",
        upload_folder,
    )

    monkeypatch.setattr(
        "app.main.organization",
        Organization(),
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "create_workspace"
        ),
        lambda **kwargs: {
            "workspace_id": "test-existing-coa-workspace",
            "status": "ready",
            "paths": {
                "financial_model": str(
                    financial_model_folder
                ),
                "ai_knowledge": str(
                    ai_knowledge_folder
                ),
            },
        },
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "save_upload"
        ),
        lambda **kwargs: {
            "status": "saved",
        },
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "General Ledger"

    sheet.append(
        [
            "Posting Date",
            "G/L Account No.",
            "G/L Account Name",
            "Debit Amount (LCY)",
            "Credit Amount (LCY)",
        ]
    )

    sheet.append(
        [
            "2026-01-15",
            "610001",
            "GL Salaries Name",
            100.0,
            0.0,
        ]
    )

    sheet.append(
        [
            "2026-01-15",
            "250001",
            "GL Asset Name",
            0.0,
            100.0,
        ]
    )

    workbook_bytes = BytesIO()

    workbook.save(
        workbook_bytes
    )

    workbook_bytes.seek(0)

    response = client.post(
        "/upload",
        files={
            "file": (
                "gl_existing_coa.xlsx",
                workbook_bytes.getvalue(),
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
            ),
        },
        data={
            "organisation_id": "existing-coa-test",
            "organisation_name": "Existing COA Test",
            "base_currency": "USD",
        },
    )

    assert response.status_code == 200

    with dim_account_file.open(
        "r",
        encoding="utf-8",
    ) as file:
        persisted_accounts = json.load(file)

    assert persisted_accounts == existing_accounts 

def test_gl_only_upload_rejects_missing_account_numbers(
    tmp_path,
    monkeypatch,
):
    upload_folder = (
        Path(tmp_path)
        / "uploads"
    )

    financial_model_folder = (
        Path(tmp_path)
        / "financial_model"
    )

    ai_knowledge_folder = (
        Path(tmp_path)
        / "ai_knowledge"
    )

    upload_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    financial_model_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    ai_knowledge_folder.mkdir(
        parents=True,
        exist_ok=True,
    )

    monkeypatch.setattr(
        "app.main.UPLOAD_FOLDER",
        upload_folder,
    )

    monkeypatch.setattr(
        "app.main.organization",
        Organization(),
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "create_workspace"
        ),
        lambda **kwargs: {
            "workspace_id": "test-gl-missing-account",
            "status": "ready",
            "paths": {
                "financial_model": str(
                    financial_model_folder
                ),
                "ai_knowledge": str(
                    ai_knowledge_folder
                ),
            },
        },
    )

    monkeypatch.setattr(
        (
            "app.main.workspace_service."
            "save_upload"
        ),
        lambda **kwargs: {
            "status": "saved",
        },
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "General Ledger"

    sheet.append(
        [
            "Posting Date",
            "G/L Account No.",
            "G/L Account Name",
            "Debit Amount (LCY)",
            "Credit Amount (LCY)",
        ]
    )

    sheet.append(
        [
            "2026-01-15",
            None,
            "Missing Account",
            100.0,
            0.0,
        ]
    )

    workbook_bytes = BytesIO()

    workbook.save(
        workbook_bytes
    )

    workbook_bytes.seek(0)

    response = client.post(
        "/upload",
        files={
            "file": (
                "gl_missing_accounts.xlsx",
                workbook_bytes.getvalue(),
                (
                    "application/vnd.openxmlformats-"
                    "officedocument.spreadsheetml.sheet"
                ),
            ),
        },
        data={
            "organisation_id": "gl-missing-account",
            "organisation_name": "GL Missing Account",
            "base_currency": "USD",
        },
    )

    assert response.status_code == 400

    assert (
        response.json()["detail"]
        == (
            "The General Ledger does not contain "
            "any usable account numbers."
        )
    )
           